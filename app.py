from flask import Flask, render_template, request, redirect, url_for, flash, send_file, jsonify
import sqlite3
import pandas as pd
import re
import calendar
import json
import io
import os
import warnings
from datetime import date, datetime
from parser import ingest_attendance, ingest_calibrations, find_attendance_header_and_period, ingest_rejections
from pptx_generator import generate_monthly_deck
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

app = Flask(__name__)
app.secret_key = "super_secret_performance_key"
DB_PATH = "performance_tracker.db"

TEMP_UPLOAD_DIR = "temp_uploads"
os.makedirs(TEMP_UPLOAD_DIR, exist_ok=True)

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def parse_iso_date(date_val):
    if pd.isna(date_val) or not date_val:
        return None
    try:
        ts = pd.to_datetime(str(date_val).strip(), errors="coerce")
        return None if pd.isna(ts) else ts.date()
    except Exception:
        return None

def normalize_str(s):
    return re.sub(r'\s+', ' ', str(s).strip().replace('"', '')).lower()

def is_working_day_status(raw_status):
    if not raw_status:
        return False
    s = str(raw_status).strip().upper()
    if s in ["NW", "AL", "MC", "PH", "PL", "CP", "UPL", "ABS", "RES", "-", ""]:
        return False
    if "|" in s:
        s = s.split("|")[0].strip()
    try:
        return float(s) > 0
    except ValueError:
        return False

def format_display_date(d_val):
    d = parse_iso_date(d_val)
    return d.strftime("%d %b %Y").lstrip("0") if d else "N/A"

def calculate_tenure(start_val, end_val=None):
    s = parse_iso_date(start_val)
    if not s:
        return None
    e = parse_iso_date(end_val) or date.today()
    days = (e - s).days
    if days < 0:
        return "0 mos"
    years = days // 365
    months = (days % 365) // 30
    if years > 0 and months > 0:
        return f"{years} yr{'s' if years > 1 else ''} {months} mo{'s' if months > 1 else ''}"
    elif years > 0:
        return f"{years} yr{'s' if years > 1 else ''}"
    return f"{max(1, months)} mo{'s' if months > 1 else ''}"

def get_keysight_fiscal_info(year, month):
    if month in [11, 12]:
        q, fy = "Q1", year + 1
    elif month == 1:
        q, fy = "Q1", year
    elif month in [2, 3, 4]:
        q, fy = "Q2", year
    elif month in [5, 6, 7]:
        q, fy = "Q3", year
    else:
        q, fy = "Q4", year
    return q, fy, f"{q} FY'{str(fy)[-2:]}"

def migrate_and_clean_database():
    conn = get_db()
    cur = conn.cursor()
    
    cur.execute("PRAGMA table_info(stations)")
    st_cols = [r["name"] for r in cur.fetchall()]
    if "sort_order" not in st_cols:
        cur.execute("ALTER TABLE stations ADD COLUMN sort_order INTEGER DEFAULT 99")
    if "target_pi" not in st_cols:
        cur.execute("ALTER TABLE stations ADD COLUMN target_pi REAL DEFAULT 3.0")
    if "is_active" not in st_cols:
        cur.execute("ALTER TABLE stations ADD COLUMN is_active BOOLEAN DEFAULT 1")
    if "group_name" not in st_cols:
        cur.execute("ALTER TABLE stations ADD COLUMN group_name TEXT")
        cur.execute("UPDATE stations SET group_name = 'CSG' WHERE UPPER(TRIM(name)) LIKE '%SIGNAL%' OR UPPER(TRIM(name)) LIKE '%MXG%'")
        cur.execute("UPDATE stations SET group_name = 'EiSG+Scope' WHERE UPPER(TRIM(name)) LIKE '%PNA%'")
        cur.execute("UPDATE stations SET group_name = 'STDS LAB' WHERE UPPER(TRIM(name)) LIKE '%STDS%'")
        cur.execute("UPDATE stations SET group_name = 'Legacy+Modular+PMPS (A)+MMA+ Others' WHERE UPPER(TRIM(name)) LIKE '%POWER%' OR UPPER(TRIM(name)) LIKE '%ZCAL%'")

    cur.execute("UPDATE stations SET name = 'Signal Analyzer' WHERE name = 'SA'")
    cur.execute("UPDATE stations SET is_active = 0 WHERE name LIKE '%MECH%'")

    cur.execute("""
    CREATE TABLE IF NOT EXISTS engineer_station_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        engineer_id INTEGER NOT NULL,
        station_id INTEGER NOT NULL,
        effective_year INTEGER NOT NULL,
        effective_month INTEGER NOT NULL,
        FOREIGN KEY (engineer_id) REFERENCES engineers (id),
        FOREIGN KEY (station_id) REFERENCES stations (id),
        UNIQUE(engineer_id, effective_year, effective_month)
    );
    """)

    # One-time safe historical repair
    mxg_st = cur.execute("SELECT id FROM stations WHERE name LIKE '%MXG%'").fetchone()
    if mxg_st:
        cur.execute("""
            INSERT OR IGNORE INTO engineer_station_history (engineer_id, station_id, effective_year, effective_month)
            SELECT h.engineer_id, ?, 2024, 11
            FROM engineer_station_history h
            JOIN stations s ON h.station_id = s.id
            WHERE s.name LIKE '%POWER SENSOR%'
              AND NOT EXISTS (
                  SELECT 1 FROM engineer_station_history h2 
                  WHERE h2.engineer_id = h.engineer_id AND (h2.effective_year < 2026 OR (h2.effective_year = 2026 AND h2.effective_month < 9))
              )
        """, (mxg_st[0],))

    cur.execute("""
    CREATE TABLE IF NOT EXISTS ignored_technicians (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        reason TEXT DEFAULT 'FTE / Untracked Tech'
    );
    """)

    cur.execute("""
    CREATE TABLE IF NOT EXISTS rejection_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        service_id TEXT, model_id TEXT, serial_no TEXT, cal_by TEXT,
        reject_date TEXT, reject_year INTEGER, reject_month INTEGER,
        failure_reason TEXT, remark TEXT
    );
    """)

    cur.execute("""
    CREATE TABLE IF NOT EXISTS model_groups (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        model_id TEXT UNIQUE NOT NULL,
        group_name TEXT NOT NULL
    );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_model_groups_model ON model_groups(model_id);")

    cur.execute("PRAGMA table_info(engineers)")
    cols = [r["name"] for r in cur.fetchall()]
    if "joined_date" not in cols:
        cur.execute("ALTER TABLE engineers ADD COLUMN joined_date TEXT")
    if "activation_date" not in cols:
        cur.execute("ALTER TABLE engineers ADD COLUMN activation_date TEXT")
    if "resigned_date" not in cols:
        cur.execute("ALTER TABLE engineers ADD COLUMN resigned_date TEXT")
    if "custom_target_pi" not in cols:
        cur.execute("ALTER TABLE engineers ADD COLUMN custom_target_pi REAL")
    if "display_name" not in cols:
        cur.execute("ALTER TABLE engineers ADD COLUMN display_name TEXT")
    if "track_performance" not in cols:
        cur.execute("ALTER TABLE engineers ADD COLUMN track_performance BOOLEAN DEFAULT 1")

    cur.execute("UPDATE engineers SET track_performance = 1 WHERE track_performance IS NULL")

    cur.execute("PRAGMA table_info(calibrations)")
    cal_cols = [r["name"] for r in cur.fetchall()]
    if "sender_dept" not in cal_cols:
        cur.execute("ALTER TABLE calibrations ADD COLUMN sender_dept TEXT")

    cur.execute("""
    CREATE TABLE IF NOT EXISTS kpi_configs (
        kpi_key TEXT PRIMARY KEY,
        kpi_name TEXT NOT NULL,
        target_label TEXT DEFAULT 'Q2 FY''26 Target',
        target_value REAL NOT NULL,
        unit TEXT DEFAULT '%'
    );
    """)

    cur.execute("""
    CREATE TABLE IF NOT EXISTS system_settings (
        setting_key TEXT PRIMARY KEY,
        setting_val TEXT NOT NULL
    );
    """)
    cur.execute("INSERT OR IGNORE INTO system_settings (setting_key, setting_val) VALUES ('target_headcount', '23')")
    cur.execute("UPDATE kpi_configs SET kpi_name = 'KGSO External CRD' WHERE kpi_key = 'kgso_ext'")

    default_kpis = [
        ("ete_priority", "ETE Priority Management System", "Q2 FY'26 Target", 90.0, "%"),
        ("ete_crd_tat", "Electrical ETE CRD TAT", "Q2 FY'26 Target", 85.0, "%"),
        ("trade_unit", "Trade Unit (CRD)", "Q2 FY'26 Target", 90.0, "%"),
        ("krs_unit", "KRS Unit (CRD)", "Q2 FY'26 Target", 92.0, "%"),
        ("kgso_ext", "KGSO External CRD", "Q2 FY'26 Target", 95.0, "%"),
        ("pi", "Productivity Index", "Q2 FY'26 Target", 3.1, "units")
    ]
    for k_key, k_name, k_label, k_val, k_unit in default_kpis:
        cur.execute("""
            INSERT OR IGNORE INTO kpi_configs (kpi_key, kpi_name, target_label, target_value, unit)
            VALUES (?, ?, ?, ?, ?)
        """, (k_key, k_name, k_label, k_val, k_unit))

    cur.execute("""
    CREATE TABLE IF NOT EXISTS quality_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        engineer_id INTEGER NOT NULL,
        year INTEGER NOT NULL,
        month INTEGER NOT NULL,
        mistakes INTEGER DEFAULT 0,
        FOREIGN KEY (engineer_id) REFERENCES engineers (id),
        UNIQUE(engineer_id, year, month)
    );
    """)

    cur.execute("CREATE INDEX IF NOT EXISTS idx_cal_period_user ON calibrations(cal_year, cal_month, cal_by);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_att_period_name ON attendance(year, month, formal_name);")

    conn.commit()
    conn.close()

migrate_and_clean_database()

TAT_CATEGORIES = [
    {"key": "CRD_ETE", "label": "ETE", "weight": 15.0},
    {"key": "CRD_NEW", "label": "NEW", "weight": 35.0},
    {"key": "PSC_EXT", "label": "PSC_EXT", "weight": 35.0},
    {"key": "PSC_INT", "label": "PSC_INT", "weight": 15.0}
]

def resolve_engineer_station(conn, engineer_id, fallback_station_id, rep_year, rep_month):
    row = conn.execute("""
        SELECT h.station_id, s.name as station_name, s.sort_order, s.target_pi, s.is_active
        FROM engineer_station_history h
        JOIN stations s ON h.station_id = s.id
        WHERE h.engineer_id = ? 
          AND (h.effective_year < ? OR (h.effective_year = ? AND h.effective_month <= ?))
        ORDER BY h.effective_year DESC, h.effective_month DESC
        LIMIT 1
    """, (engineer_id, rep_year, rep_year, rep_month)).fetchone()

    if row:
        return row["station_id"], row["station_name"], row["sort_order"], row["target_pi"], row["is_active"]
    
    fb = conn.execute("""
        SELECT s.id as station_id, s.name as station_name, s.sort_order, s.target_pi, s.is_active
        FROM stations s WHERE s.id = ?
    """, (fallback_station_id,)).fetchone()
    
    if fb:
        return fb["station_id"], fb["station_name"], fb["sort_order"], fb["target_pi"], fb["is_active"]
    return None, "Unassigned", 99, 3.0, 1

@app.route("/", strict_slashes=False)
def dashboard():
    conn = get_db()
    query_periods = """
        SELECT DISTINCT cal_year as yr, cal_month as mo FROM calibrations WHERE cal_year IS NOT NULL
        UNION
        SELECT DISTINCT year as yr, month as mo FROM attendance WHERE year IS NOT NULL
        ORDER BY yr DESC, mo DESC
    """
    raw_periods = conn.execute(query_periods).fetchall()

    available_periods = []
    for r in raw_periods:
        if r["yr"] and r["mo"]:
            val_str = f"{r['yr']}-{int(r['mo']):02d}"
            label_str = f"{calendar.month_name[int(r['mo'])]} {r['yr']}"
            available_periods.append({"val": val_str, "label": label_str, "year": int(r["yr"]), "month": int(r["mo"])})

    req_period = request.args.get("period")
    selected_period = next((p for p in available_periods if p["val"] == req_period), None)
    if not selected_period:
        selected_period = available_periods[0] if available_periods else {"val": "2026-08", "label": "August 2026", "year": 2026, "month": 8}

    rep_year = selected_period["year"]
    rep_month = selected_period["month"]

    _, num_days = calendar.monthrange(rep_year, rep_month)
    month_start, month_end = date(rep_year, rep_month, 1), date(rep_year, rep_month, num_days)
    days_list = list(range(1, num_days + 1))
    
    df_att_raw = pd.read_sql("SELECT day, raw_status FROM attendance WHERE month = ? AND year = ?", conn, params=[rep_month, rep_year])
    ph_days = set(df_att_raw[df_att_raw["raw_status"].astype(str).str.strip().str.upper() == "PH"]["day"].astype(int).tolist()) if not df_att_raw.empty else set()

    days_info = []
    for d in days_list:
        cur_date = date(rep_year, rep_month, d)
        days_info.append({
            "day_num": d, "day_str": f"{d:02d}", "weekday": cur_date.strftime("%a"),
            "is_weekend": cur_date.weekday() >= 5, "is_ph": (d in ph_days)
        })

    report_period = selected_period["label"]
    df_mistakes = pd.read_sql("SELECT engineer_id, mistakes FROM quality_logs WHERE year = ? AND month = ?", conn, params=[rep_year, rep_month])
    mistakes_map = dict(zip(df_mistakes["engineer_id"], df_mistakes["mistakes"])) if not df_mistakes.empty else {}

    all_engineers = conn.execute("SELECT * FROM engineers WHERE track_performance = 1 ORDER BY formal_name").fetchall()
    roster_usernames = set([e["username"].strip().lower() for e in all_engineers])
    muted_usernames = set([m["username"].strip().lower() for m in conn.execute("SELECT username FROM ignored_technicians").fetchall()])

    df_cal_counts_all = pd.read_sql("SELECT cal_by, COUNT(*) as units FROM calibrations WHERE cal_month = ? AND cal_year = ? GROUP BY cal_by", conn, params=[rep_month, rep_year])
    user_cal_totals = {}
    unmapped_techs = []
    if not df_cal_counts_all.empty:
        for _, row in df_cal_counts_all.iterrows():
            c_by, u_cnt = str(row["cal_by"]).strip().lower(), int(row["units"])
            user_cal_totals[c_by] = u_cnt
            if c_by not in roster_usernames and c_by not in muted_usernames:
                unmapped_techs.append({"username": str(row["cal_by"]).strip(), "units": u_cnt})

    eligible_staff = []
    for e in all_engineers:
        u_clean = str(e["username"]).strip().lower()
        act_d, res_d = parse_iso_date(e["activation_date"]), parse_iso_date(e["resigned_date"])
        units_done = user_cal_totals.get(u_clean, 0)

        if not act_d or act_d > month_end: continue
        if res_d and res_d < month_start and units_done == 0: continue

        st_id, st_name, st_order, st_target, st_active = resolve_engineer_station(conn, e["id"], e["station_id"], rep_year, rep_month)
        if not st_active: continue

        ind_target = e["custom_target_pi"] if e["custom_target_pi"] is not None else st_target
        e_dict = dict(e)
        e_dict.update({"station_id": st_id, "station_name": st_name, "station_sort_order": st_order, "station_target_pi": st_target, "target_pi": ind_target})
        eligible_staff.append(e_dict)

    if not eligible_staff:
        conn.close()
        return render_template("dashboard.html", empty=True, days_info=days_info, summary=[], station_metrics={}, warnings=[], report_period=report_period, tat_categories=TAT_CATEGORIES, available_periods=available_periods, selected_period=selected_period["val"], rep_year=rep_year, rep_month=rep_month, unmapped_techs=unmapped_techs, dept_total_units=0, dept_avg_pi=0.0)

    eligible_staff.sort(key=lambda x: (x["station_sort_order"], x["station_name"], x["formal_name"]))
    usernames = [s["username"] for s in eligible_staff]

    q_cal_all = f"SELECT cal_by, cal_day, category, mos FROM calibrations WHERE cal_by IN ({','.join(['?']*len(usernames))}) AND cal_month = ? AND cal_year = ?"
    df_cal_all = pd.read_sql(q_cal_all, conn, params=usernames + [rep_month, rep_year])

    if not df_cal_all.empty:
        df_cal_counts = df_cal_all.groupby(["cal_by", "cal_day"]).size().reset_index(name="units")
        pivot_cal = df_cal_counts.pivot(index="cal_by", columns="cal_day", values="units").reindex(index=usernames, columns=days_list).fillna(0).astype(int)
    else:
        pivot_cal = pd.DataFrame(0, index=usernames, columns=days_list)

    df_att_all = pd.read_sql("SELECT formal_name, day, raw_status, capped_hours FROM attendance WHERE month = ? AND year = ?", conn, params=[rep_month, rep_year])
    if not df_att_all.empty:
        df_att_all["norm_name"] = df_att_all["formal_name"].apply(normalize_str)

    individual_summary, warnings = [], []
    station_metrics = {}

    for s in eligible_staff:
        eng_id, u, fn, dn = s["id"], s["username"], s["formal_name"], s["display_name"] or s["formal_name"].split()[0]
        fn_norm, st, st_order, ind_target, shift = normalize_str(fn), s["station_name"], s["station_sort_order"], s["target_pi"], s["shift_hours"]
        res_d, act_d = parse_iso_date(s["resigned_date"]), parse_iso_date(s["activation_date"])
        
        lifecycle_badge = None
        if res_d and res_d <= month_end:
            lifecycle_badge = {"text": f"Resigned on {res_d.strftime('%d/%m/%Y')}", "class": "danger"}
        elif act_d and month_start <= act_d <= month_end:
            lifecycle_badge = {"text": f"Activated on {act_d.strftime('%d/%m/%Y')}", "class": "info text-dark"}

        units = int(pivot_cal.loc[u].sum()) if u in pivot_cal.index else 0
        mistakes = mistakes_map.get(eng_id, 0)
        quality_pct = max(0.0, round(((units - mistakes) / units) * 100.0, 1)) if units > 0 else 100.0

        person_att = df_att_all[df_att_all["norm_name"] == fn_norm] if not df_att_all.empty else pd.DataFrame()
        has_attendance = not person_att.empty

        if has_attendance:
            att_map = dict(zip(person_att["day"], person_att["raw_status"]))
            hrs = float(person_att["capped_hours"].sum())
        else:
            att_map, hrs = {}, 0.0
            if not lifecycle_badge or "Resigned" not in lifecycle_badge["text"]:
                warnings.append(f"Missing Attendance Record: {fn} ({st})")

        raw_eff_days = (hrs / shift) if shift > 0 else 0.0
        eff_days = int(round(raw_eff_days))
        pi = round(units / raw_eff_days, 1) if raw_eff_days > 0 else 0.0

        if st not in station_metrics:
            station_metrics[st] = {"units": 0, "eff_days": 0.0, "headcount": 0, "sort_order": st_order, "target_units_sum": 0.0, "individual_targets": []}
        station_metrics[st]["units"] += units
        station_metrics[st]["eff_days"] += raw_eff_days
        station_metrics[st]["headcount"] += 1
        station_metrics[st]["target_units_sum"] += (raw_eff_days * ind_target)
        station_metrics[st]["individual_targets"].append(ind_target)

        raw_out_list = [int(pivot_cal.loc[u].get(d, 0)) for d in days_list] if u in pivot_cal.index else [0]*num_days
        raw_att_list = [str(att_map.get(d, "-")) for d in days_list]
        working_flags = [is_working_day_status(raw_att_list[i]) for i in range(num_days)]

        adj_out_list = list(raw_out_list)
        carry = 0
        for i in range(num_days):
            if not working_flags[i]:
                if adj_out_list[i] > 0:
                    carry += adj_out_list[i]
                    adj_out_list[i] = 0
            else:
                if carry > 0:
                    adj_out_list[i] += carry
                    carry = 0
        if carry > 0:
            for i in range(num_days - 1, -1, -1):
                if working_flags[i]:
                    adj_out_list[i] += carry
                    break

        formatted_att_list = ["" if (raw_att_list[i] in ["NW", "-", "0", "0.0", "", "None"]) and not working_flags[i] else raw_att_list[i] for i in range(num_days)]
        formatted_out_list = []
        for i in range(num_days):
            out_val, is_work = adj_out_list[i], working_flags[i]
            formatted_out_list.append({
                "val": ("" if out_val == 0 else str(out_val)) if not is_work else str(out_val),
                "is_zero_work": (is_work and out_val == 0),
                "has_val": out_val > 0
            })

        user_cals = df_cal_all[df_cal_all["cal_by"] == u] if not df_cal_all.empty else pd.DataFrame()
        tat_scores = {}
        weighted_num, active_weights = 0.0, 0.0

        for cat in TAT_CATEGORIES:
            cat_key, cat_label, weight = cat["key"], cat["label"], cat["weight"]
            cat_rows = user_cals[user_cals["category"].str.strip().str.upper() == cat_key.upper()] if not user_cals.empty else pd.DataFrame()
            total_cat = len(cat_rows)

            if total_cat == 0:
                tat_scores[cat_label] = {"val": "N/A", "numeric": None, "total": 0, "met": 0}
            else:
                met_count = int((cat_rows["mos"] == 1).sum())
                pct = int(round((met_count / total_cat) * 100.0))
                tat_scores[cat_label] = {"val": f"{pct}%", "numeric": pct, "total": total_cat, "met": met_count}
                weighted_num += pct * (weight / 100.0)
                active_weights += (weight / 100.0)

        overall_tat_val = int(round(weighted_num / active_weights)) if active_weights > 0 else None
        overall_tat_str = f"{overall_tat_val}%" if overall_tat_val is not None else "N/A"

        individual_summary.append({
            "username": u, "formal_name": fn, "display_name": dn, "station": st, "target_pi": ind_target,
            "shift_hours": shift, "total_units": units, "mistakes": mistakes, "quality_pct": quality_pct,
            "total_hours": round(hrs, 1), "total_hours_int": int(round(hrs)), "effective_days": eff_days,
            "pi": pi, "pi_met": pi >= ind_target, "has_attendance": has_attendance,
            "daily_output": formatted_out_list, "daily_attendance": formatted_att_list,
            "tat_scores": tat_scores, "overall_tat_val": overall_tat_val, "overall_tat_str": overall_tat_str,
            "lifecycle_badge": lifecycle_badge
        })

    sorted_station_metrics = dict(sorted(station_metrics.items(), key=lambda item: (item[1]["sort_order"], item[0])))
    for st, m in sorted_station_metrics.items():
        m["pi"] = round(m["units"] / m["eff_days"], 1) if m["eff_days"] > 0 else 0.0
        m["target_pi"] = round(m["target_units_sum"] / m["eff_days"], 1) if m["eff_days"] > 0 else (round(sum(m["individual_targets"]) / len(m["individual_targets"]), 1) if m["individual_targets"] else 3.0)
        m["pi_met"] = m["pi"] >= m["target_pi"]

    dept_total_units = sum(m["units"] for m in sorted_station_metrics.values())
    num_stations = len(sorted_station_metrics)
    dept_avg_pi = round(sum(m["pi"] for m in sorted_station_metrics.values()) / num_stations, 1) if num_stations > 0 else 0.0

    conn.close()
    return render_template(
        "dashboard.html", days_info=days_info, days_list=days_list, summary=individual_summary,
        station_metrics=sorted_station_metrics, warnings=warnings, report_period=report_period,
        tat_categories=TAT_CATEGORIES, available_periods=available_periods, selected_period=selected_period["val"],
        rep_year=rep_year, rep_month=rep_month, unmapped_techs=unmapped_techs,
        dept_total_units=dept_total_units, dept_avg_pi=dept_avg_pi
    )

@app.route("/api/update_cell", methods=["POST"])
def update_cell():
    data = request.get_json()
    if not data: return jsonify({"success": False, "error": "No data"}), 400

    c_type, username, formal_name = data.get("type"), data.get("username"), data.get("formal_name")
    day, year, month, val = int(data.get("day")), int(data.get("year")), int(data.get("month")), str(data.get("value", "")).strip()

    conn = get_db()
    cur = conn.cursor()
    try:
        if c_type == "hours":
            eng = cur.execute("SELECT shift_hours FROM engineers WHERE username = ?", (username,)).fetchone()
            shift = eng[0] if eng else 8.0
            if val.upper() in ["NW", "AL", "MC", "PH", "PL", "CP", "UPL", "ABS", "RES"]:
                raw_st, act_h, cap_h = val.upper(), 0.0, 0.0
            else:
                try: h = float(val) if val else 0.0
                except ValueError: h = 0.0
                act_h = h
                cap_h = min(h, 10.0) if shift == 8.0 else h
                raw_st = f"{cap_h:.1f}".rstrip('0').rstrip('.') if cap_h > 0 else ""

            row = cur.execute("SELECT id FROM attendance WHERE LOWER(TRIM(formal_name)) = LOWER(TRIM(?)) AND day = ? AND month = ? AND year = ?", (formal_name, day, month, year)).fetchone()
            if row:
                cur.execute("UPDATE attendance SET raw_status = ?, actual_hours = ?, capped_hours = ? WHERE id = ?", (raw_st, act_h, cap_h, row[0]))
            else:
                cur.execute("INSERT INTO attendance (formal_name, day, month, year, raw_status, actual_hours, capped_hours) VALUES (?, ?, ?, ?, ?, ?, ?)", (formal_name, day, month, year, raw_st, act_h, cap_h))

        elif c_type == "output":
            target_count = max(0, int(val)) if val.isdigit() else 0
            cur_rows = cur.execute("SELECT id FROM calibrations WHERE cal_by = ? AND cal_day = ? AND cal_month = ? AND cal_year = ?", (username, day, month, year)).fetchall()
            current_count = len(cur_rows)
            if target_count > current_count:
                for _ in range(target_count - current_count):
                    cur.execute("INSERT INTO calibrations (cal_by, cal_day, cal_month, cal_year, cal_date, category, mos) VALUES (?, ?, ?, ?, ?, 'CRD_ETE', 1)", (username, day, month, year, f"{year}-{month:02d}-{day:02d}"))
            elif target_count < current_count:
                ids_to_del = [r[0] for r in cur_rows[:(current_count - target_count)]]
                cur.execute(f"DELETE FROM calibrations WHERE id IN ({','.join(['?']*len(ids_to_del))})", ids_to_del)

        conn.commit()
        conn.close()
        return jsonify({"success": True})
    except Exception as e:
        conn.close()
        return jsonify({"success": False, "error": str(e)}), 500

@app.route("/mute_technician")
def mute_technician():
    username = request.args.get("username", "").strip()
    period = request.args.get("period", "2026-08")
    redir = request.args.get("redirect", "dashboard")
    if username:
        conn = get_db()
        conn.execute("INSERT OR IGNORE INTO ignored_technicians (username, reason) VALUES (?, 'FTE / Untracked Tech')", (username,))
        conn.commit()
        conn.close()
        flash(f"Technician '{username}' muted.", "info")
    return redirect(url_for("upload") if redir == "preflight" else url_for("dashboard", period=period))

@app.route("/unmute_technician/<int:m_id>")
def unmute_technician(m_id):
    conn = get_db()
    conn.execute("DELETE FROM ignored_technicians WHERE id = ?", (m_id,))
    conn.commit()
    conn.close()
    flash("Technician unmuted successfully.", "success")
    return redirect(url_for("team"))

@app.route("/export_excel", strict_slashes=False)
def export_excel():
    conn = get_db()
    req_period = request.args.get("period", "2026-08")
    try: y, m = [int(x) for x in req_period.split("-")]
    except ValueError: y, m = 2026, 8

    wb = openpyxl.Workbook()
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    emerald_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    white_font, reg_font = Font(name="Arial", size=10, bold=True, color="FFFFFF"), Font(name="Arial", size=9)
    red_font, green_font = Font(name="Arial", size=9, bold=True, color="DC3545"), Font(name="Arial", size=9, bold=True, color="198754")
    center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin', color='B0C4DE'), right=Side(style='thin', color='B0C4DE'), top=Side(style='thin', color='B0C4DE'), bottom=Side(style='thin', color='B0C4DE'))

    ws_pi = wb.active
    ws_pi.title = "P.I. & Quality Summary"
    ws_pi.append(["Station", "Engineer Name", "Total Output", "Mistakes", "Quality %", "Total Hours", "Shift", "Equiv. Days", "Target P.I.", "Actual P.I."])
    for col_idx in range(1, 11):
        c = ws_pi.cell(row=1, column=col_idx)
        c.fill, c.font, c.alignment = navy_fill, white_font, center

    all_engineers = conn.execute("SELECT e.*, s.name as station_name, s.target_pi as st_target, s.sort_order FROM engineers e LEFT JOIN stations s ON e.station_id = s.id WHERE e.track_performance = 1 AND s.is_active = 1 ORDER BY s.sort_order ASC, s.name ASC, e.formal_name ASC").fetchall()
    df_cals = pd.read_sql("SELECT cal_by, category, mos, cal_day FROM calibrations WHERE cal_year = ? AND cal_month = ?", conn, params=[y, m])
    u_map = df_cals["cal_by"].value_counts().to_dict() if not df_cals.empty else {}
    df_att = pd.read_sql("SELECT formal_name, SUM(capped_hours) as hrs FROM attendance WHERE year = ? AND month = ?", conn, params=[y, m])
    h_map = dict(zip(df_att["formal_name"].str.strip().str.lower(), df_att["hrs"])) if not df_att.empty else {}
    df_mis = pd.read_sql("SELECT engineer_id, mistakes FROM quality_logs WHERE year = ? AND month = ?", conn, params=[y, m])
    m_map = dict(zip(df_mis["engineer_id"], df_mis["mistakes"])) if not df_mis.empty else {}

    r_idx = 2
    for e in all_engineers:
        u, fn, st = e["username"], e["formal_name"], e["station_name"] or "Unassigned"
        target_pi = e["custom_target_pi"] if e["custom_target_pi"] is not None else (e["st_target"] or 3.0)
        units = u_map.get(u, 0)
        hrs = h_map.get(normalize_str(fn), 0.0)
        mistakes = m_map.get(e["id"], 0)
        q_pct = max(0.0, round(((units - mistakes) / units) * 100.0, 1)) if units > 0 else 100.0
        eff_days = int(round(hrs / e["shift_hours"])) if e["shift_hours"] > 0 else 0
        actual_pi = round(units / (hrs / e["shift_hours"]), 1) if (hrs > 0 and e["shift_hours"] > 0) else 0.0

        if units > 0 or hrs > 0:
            ws_pi.append([st, fn, units, mistakes, f"{q_pct}%", f"{round(hrs, 1)} hrs", f"{e['shift_hours']}h", eff_days, target_pi, actual_pi])
            for c_idx in range(1, 11):
                c = ws_pi.cell(row=r_idx, column=c_idx)
                c.font, c.alignment, c.border = reg_font, center, thin_border
                if c_idx == 5: c.font = green_font if q_pct >= 98.0 else red_font
                elif c_idx == 10: c.font = green_font if actual_pi >= target_pi else red_font
            r_idx += 1

    ws_tat = wb.create_sheet(title="TAT Compliance")
    ws_tat.append(["Station", "Engineer Name", "Total Units", "ETE (15%)", "NEW (35%)", "PSC_EXT (35%)", "PSC_INT (15%)", "Overall TAT %"])
    for col_idx in range(1, 9):
        c = ws_tat.cell(row=1, column=col_idx)
        c.fill, c.font, c.alignment = emerald_fill, white_font, center

    r_tat_idx = 2
    for e in all_engineers:
        u, fn, st = e["username"], e["formal_name"], e["station_name"] or "Unassigned"
        u_cals = df_cals[df_cals["cal_by"] == u] if not df_cals.empty else pd.DataFrame()
        tot_units = len(u_cals)
        cat_scores = {}
        w_num, act_w = 0.0, 0.0
        for cat in TAT_CATEGORIES:
            k, w, lbl = cat["key"], cat["weight"], cat["label"].lower()
            c_rows = u_cals[u_cals["category"].str.strip().str.upper() == k] if not u_cals.empty else pd.DataFrame()
            if len(c_rows) > 0:
                pct = int(round(((c_rows["mos"] == 1).sum() / len(c_rows)) * 100.0))
                cat_scores[lbl] = f"{pct}%"
                w_num += pct * (w / 100.0)
                act_w += (w / 100.0)
            else: cat_scores[lbl] = "N/A"

        tat_val = int(round(w_num / act_w)) if act_w > 0 else None
        tat_str = f"{tat_val}%" if tat_val is not None else "N/A"

        if tot_units > 0:
            ws_tat.append([st, fn, tot_units, cat_scores.get("ete", "N/A"), cat_scores.get("new", "N/A"), cat_scores.get("psc_ext", "N/A"), cat_scores.get("psc_int", "N/A"), tat_str])
            for c_idx in range(1, 9):
                c = ws_tat.cell(row=r_tat_idx, column=c_idx)
                c.font, c.alignment, c.border = reg_font, center, thin_border
                if c_idx == 8 and tat_val is not None: c.font = green_font if tat_val >= 90 else red_font
            r_tat_idx += 1

    for sheet in [ws_pi, ws_tat]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 3, 12)

    conn.close()
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    return send_file(excel_stream, as_attachment=True, download_name=f"Performance_Report_{calendar.month_abbr[m]}_{y}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/export_engineer_calibrations/<int:eng_id>", strict_slashes=False)
def export_engineer_calibrations(eng_id):
    conn = get_db()
    eng = conn.execute("SELECT formal_name, username FROM engineers WHERE id = ?", (eng_id,)).fetchone()
    if not eng:
        conn.close()
        flash("Engineer not found.", "danger")
        return redirect(url_for("history"))

    u, fn = eng["username"], eng["formal_name"]
    df_cals = pd.read_sql("""
        SELECT service_id as [Service ID], model_id as [Model ID], serial_no as [Serial No], 
               cal_date as [Cal Date], cal_day as [Day], cal_month as [Month], cal_year as [Year], 
               category as [Category], CASE WHEN mos = 1 THEN 'Met' ELSE 'Late' END as [TAT Status (MOS)], sender_dept as [Sender Dept]
        FROM calibrations WHERE cal_by = ? ORDER BY cal_date DESC
    """, conn, params=[u])
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calibrated Units"
    navy_fill, white_font, reg_font = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"), Font(name="Arial", size=10, bold=True, color="FFFFFF"), Font(name="Arial", size=9)
    center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin', color='B0C4DE'), right=Side(style='thin', color='B0C4DE'), top=Side(style='thin', color='B0C4DE'), bottom=Side(style='thin', color='B0C4DE'))

    ws.append(list(df_cals.columns))
    for c_i in range(1, len(df_cals.columns) + 1):
        cell = ws.cell(row=1, column=c_i)
        cell.fill, cell.font, cell.alignment = navy_fill, white_font, center

    for r_i, row in enumerate(df_cals.itertuples(index=False), start=2):
        ws.append(list(row))
        for c_i in range(1, len(df_cals.columns) + 1):
            cell = ws.cell(row=r_i, column=c_i)
            cell.font, cell.alignment, cell.border = reg_font, center, thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 3, 12)

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    clean_fn = re.sub(r'[^a-zA-Z0-9_]', '_', fn)
    return send_file(excel_stream, as_attachment=True, download_name=f"Calibrations_{clean_fn}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# =============================================================================
# UNIFIED QUARTERLY ENGINE (WEIGHTED TAT, QOQ DELTAS, OVERALL ROW)
# =============================================================================
def calculate_quarterly_engine(conn, selected_quarter):
    periods_db = conn.execute("SELECT DISTINCT cal_year as yr, cal_month as mo FROM calibrations WHERE cal_year IS NOT NULL AND cal_month IS NOT NULL ORDER BY yr ASC, mo ASC").fetchall()
    fiscal_quarters_map = {}
    for r in periods_db:
        cal_y, m = int(r["yr"]), int(r["mo"])
        q_label, fiscal_yr, _ = get_keysight_fiscal_info(cal_y, m)
        q_key = f"{q_label} FY'{str(fiscal_yr)[-2:]}"
        if q_key not in fiscal_quarters_map:
            fiscal_quarters_map[q_key] = {"key": q_key, "months": [], "fiscal_yr": fiscal_yr, "q_num": int(q_label.replace("Q", ""))}
        fiscal_quarters_map[q_key]["months"].append((cal_y, m))

    if not fiscal_quarters_map:
        return [], {}, [], [], [], []

    sorted_quarters = list(fiscal_quarters_map.values())
    sorted_quarters.sort(key=lambda x: (x["fiscal_yr"], x["q_num"]))

    # Support for Full Fiscal Year (e.g. FULL_FY_2026)
    if selected_quarter and selected_quarter.startswith("FULL_FY_"):
        target_fy = int(selected_quarter.replace("FULL_FY_", ""))
        q_key = selected_quarter
        quarter_months = []
        prev_quarter_months = []
        for r in periods_db:
            cal_y, m = int(r["yr"]), int(r["mo"])
            _, fy, _ = get_keysight_fiscal_info(cal_y, m)
            if fy == target_fy:
                quarter_months.append((cal_y, m))
            elif fy == target_fy - 1:
                prev_quarter_months.append((cal_y, m))
    else:
        q_key = selected_quarter if selected_quarter in fiscal_quarters_map else sorted_quarters[-1]["key"]
        quarter_months = fiscal_quarters_map[q_key]["months"]
        current_idx = next((i for i, q in enumerate(sorted_quarters) if q["key"] == q_key), -1)
        prev_quarter_months = sorted_quarters[current_idx - 1]["months"] if current_idx > 0 else []

    all_engineers = conn.execute("""
        SELECT e.id, e.formal_name, e.username, e.station_id, e.shift_hours, e.custom_target_pi, 
               e.activation_date, e.resigned_date, s.name as station_name, s.target_pi as st_target, s.group_name 
        FROM engineers e LEFT JOIN stations s ON e.station_id = s.id 
        WHERE (e.track_performance = 1 OR e.track_performance IS NULL)
    """).fetchall()

    muted_rows = conn.execute("SELECT username FROM ignored_technicians").fetchall()
    muted_users = set([m["username"].strip().lower() for m in muted_rows])
    target_groups = ["CSG", "EiSG+Scope", "STDS LAB", "Legacy+Modular+PMPS (A)+MMA+ Others", "MEC Lab"]

    # Helper function to evaluate quarter metrics
    def evaluate_quarter(months_list):
        if not months_list: return {}, 0, 0.0, 0, 0.0, pd.DataFrame(), pd.DataFrame(), pd.DataFrame()
        df_c = pd.DataFrame()
        for (cy, cm) in months_list:
            df_m = pd.read_sql("SELECT service_id, model_id, serial_no, cal_date, cal_day, cal_month, cal_year, cal_by, category, mos FROM calibrations WHERE cal_year = ? AND cal_month = ?", conn, params=[cy, cm])
            if not df_m.empty: df_c = pd.concat([df_c, df_m], ignore_index=True)

        if not df_c.empty:
            df_c["clean_by"] = df_c["cal_by"].astype(str).str.strip().str.lower()
            df_c = df_c[~df_c["clean_by"].isin(muted_users)]

        df_a = pd.DataFrame()
        for (cy, cm) in months_list:
            df_m = pd.read_sql("SELECT formal_name, capped_hours, year, month FROM attendance WHERE year = ? AND month = ?", conn, params=[cy, cm])
            if not df_m.empty: df_a = pd.concat([df_a, df_m], ignore_index=True)
        if not df_a.empty: df_a["clean_name"] = df_a["formal_name"].apply(normalize_str)

        df_r = pd.DataFrame()
        for (cy, cm) in months_list:
            df_rej = pd.read_sql("SELECT service_id, model_id, serial_no, cal_by, reject_date, reject_year, reject_month, failure_reason, remark FROM rejection_logs WHERE reject_year = ? AND reject_month = ?", conn, params=[cy, cm])
            if not df_rej.empty: df_r = pd.concat([df_r, df_rej], ignore_index=True)
        if not df_r.empty: df_r["clean_by"] = df_r["cal_by"].astype(str).str.strip().str.lower()

        grp_res = {}
        for grp in target_groups:
            st_rows = conn.execute("SELECT id, name FROM stations WHERE is_active = 1 AND group_name = ?", (grp,)).fetchall()
            if not st_rows: continue
            st_ids = [s["id"] for s in st_rows]
            grp_aes = set()
            grp_days = 0.0

            for (cy, cm) in months_list:
                _, m_days = calendar.monthrange(cy, cm)
                m_start, m_end = date(cy, cm, 1), date(cy, cm, m_days)
                for e in all_engineers:
                    act_d, res_d = parse_iso_date(e["activation_date"]), parse_iso_date(e["resigned_date"])
                    if act_d and act_d > m_end: continue
                    if res_d and res_d < m_start: continue
                    st_id, _, _, _, st_active = resolve_engineer_station(conn, e["id"], e["station_id"], cy, cm)
                    if not st_active: continue
                    if st_id in st_ids:
                        u_clean = e["username"].strip().lower()
                        grp_aes.add(u_clean)
                        fn_norm = normalize_str(e["formal_name"])
                        shift = e["shift_hours"] if e["shift_hours"] > 0 else 8.0
                        hrs = df_a[(df_a["clean_name"] == fn_norm) & (df_a["year"] == cy) & (df_a["month"] == cm)]["capped_hours"].sum() if not df_a.empty else 0.0
                        grp_days += (hrs / shift)

            grp_cals = df_c[df_c["clean_by"].isin(grp_aes)] if (not df_c.empty and grp_aes) else pd.DataFrame()
            units = len(grp_cals)

            # UNIFIED WEIGHTED TAT (15/35/35/15)
            w_num, a_w = 0.0, 0.0
            for cat in TAT_CATEGORIES:
                k_key, weight = cat["key"], cat["weight"]
                c_sub = grp_cals[grp_cals["category"].str.strip().str.upper() == k_key.upper()] if not grp_cals.empty else pd.DataFrame()
                if len(c_sub) > 0:
                    pct = int(round(((c_sub["mos"] == 1).sum() / len(c_sub)) * 100.0))
                    w_num += pct * (weight / 100.0)
                    a_w += (weight / 100.0)
            weighted_tat = int(round(w_num / a_w)) if a_w > 0 else 100

            pi = round(units / grp_days, 1) if grp_days > 0 else 0.0
            grp_errors = len(df_r[df_r["clean_by"].isin(grp_aes)]) if (not df_r.empty and grp_aes) else 0
            grp_res[grp] = {"units": units, "days": grp_days, "pi": pi, "tat": weighted_tat, "errors": grp_errors}

        all_units = sum(g["units"] for g in grp_res.values())
        all_days = sum(g["days"] for g in grp_res.values())
        all_pi = round(all_units / all_days, 1) if all_days > 0 else 0.0
        all_errors = sum(g["errors"] for g in grp_res.values())

        tot_w_num, tot_a_w = 0.0, 0.0
        for cat in TAT_CATEGORIES:
            k_key, weight = cat["key"], cat["weight"]
            c_sub = df_c[df_c["category"].str.strip().str.upper() == k_key.upper()] if not df_c.empty else pd.DataFrame()
            if len(c_sub) > 0:
                pct = int(round(((c_sub["mos"] == 1).sum() / len(c_sub)) * 100.0))
                tot_w_num += pct * (weight / 100.0)
                tot_a_w += (weight / 100.0)
        tot_tat = int(round(tot_w_num / tot_a_w)) if tot_a_w > 0 else 100

        return grp_res, all_units, all_days, all_pi, tot_tat, all_errors, df_c, df_a, df_r

    cur_grp, cur_tot_u, cur_tot_d, cur_tot_pi, cur_tot_tat, cur_tot_err, df_cals_q, df_att_q, df_rejs_q = evaluate_quarter(quarter_months)
    prev_grp, prev_tot_u, prev_tot_d, prev_tot_pi, prev_tot_tat, prev_tot_err, _, _, _ = evaluate_quarter(prev_quarter_months)

    quarterly_group_data, raw_calibrations_rows, raw_labor_rows, raw_rejections_rows = [], [], [], []
    ae_station_name_map = {}

    for grp in target_groups:
        st_rows = conn.execute("SELECT id, name FROM stations WHERE is_active = 1 AND group_name = ? ORDER BY sort_order ASC, name ASC", (grp,)).fetchall()
        if not st_rows: continue

        supporting_stations = [s["name"] for s in st_rows]
        supporting_st_ids = [s["id"] for s in st_rows]
        grp_m = cur_grp.get(grp, {"units": 0, "days": 0.0, "pi": 0.0, "tat": 100, "errors": 0})
        prev_m = prev_grp.get(grp, {"pi": None, "tat": None, "errors": None})

        pi_diff = round(grp_m["pi"] - prev_m["pi"], 1) if (prev_m["pi"] is not None and prev_m["pi"] > 0) else None
        qoq_pi_str = f"+{pi_diff:.1f}" if (pi_diff and pi_diff > 0) else (f"{pi_diff:.1f}" if (pi_diff and pi_diff < 0) else "(-)")

        tat_diff = int(round(grp_m["tat"] - prev_m["tat"])) if prev_m["tat"] is not None else None
        qoq_tat_str = f"+{tat_diff}%" if (tat_diff and tat_diff > 0) else (f"{tat_diff}%" if (tat_diff and tat_diff < 0) else "(-)")

        err_diff = int(grp_m["errors"] - prev_m["errors"]) if prev_m["errors"] is not None else None
        qoq_err_str = f"+{err_diff}" if (err_diff and err_diff > 0) else (f"{err_diff}" if (err_diff and err_diff < 0) else "(-)")

        vol = grp_m["units"]
        avg_mo = int(round(vol / 3.0))
        errs = grp_m["errors"]
        qual = max(0.0, round(((vol - errs) / vol) * 100.0, 1)) if vol > 0 else 100.0

        grp_aes = set()
        for (cy, cm) in quarter_months:
            _, m_days = calendar.monthrange(cy, cm)
            m_start, m_end = date(cy, cm, 1), date(cy, cm, m_days)
            for e in all_engineers:
                act_d, res_d = parse_iso_date(e["activation_date"]), parse_iso_date(e["resigned_date"])
                if act_d and act_d > m_end: continue
                if res_d and res_d < m_start: continue
                st_id, st_name, _, _, st_active = resolve_engineer_station(conn, e["id"], e["station_id"], cy, cm)
                if not st_active: continue
                if st_id in supporting_st_ids:
                    u_clean = e["username"].strip().lower()
                    grp_aes.add(u_clean)
                    ae_station_name_map[u_clean] = st_name
                    fn_norm = normalize_str(e["formal_name"])
                    shift = e["shift_hours"] if e["shift_hours"] > 0 else 8.0
                    hrs = df_att_q[(df_att_q["clean_name"] == fn_norm) & (df_att_q["year"] == cy) & (df_att_q["month"] == cm)]["capped_hours"].sum() if not df_att_q.empty else 0.0
                    raw_labor_rows.append([e["formal_name"], e["username"], st_name, cy, cm, round(hrs, 1), shift, round(hrs/shift, 2), grp])

        if not df_cals_q.empty and grp_aes:
            for _, r in df_cals_q[df_cals_q["clean_by"].isin(grp_aes)].iterrows():
                raw_calibrations_rows.append([
                    r["service_id"], r["model_id"], r["serial_no"], r["cal_date"], 
                    r["cal_month"], r["cal_year"], r["cal_by"], r["category"], 
                    "Met" if r["mos"] == 1 else "Late", grp, ae_station_name_map.get(r["clean_by"], "Unknown")
                ])

        quarterly_group_data.append({
            "group_name": grp,
            "supporting_stations": ", ".join(supporting_stations) or "None",
            "volume": vol,
            "avg_monthly_output": avg_mo,
            "pi": grp_m["pi"],
            "qoq_pi_str": qoq_pi_str,
            "qoq_pi_val": pi_diff or 0.0,
            "tat": grp_m["tat"],
            "tat_met": grp_m["tat"] >= 90,
            "qoq_tat_str": qoq_tat_str,
            "qoq_tat_val": tat_diff or 0,
            "errors": errs,
            "qoq_err_str": qoq_err_str,
            "qoq_err_val": err_diff or 0,
            "quality": qual,
            "quality_met": qual >= 98.0
        })

    tot_pi_diff = round(cur_tot_pi - prev_tot_pi, 1) if (prev_tot_pi and prev_tot_pi > 0) else None
    tot_tat_diff = int(round(cur_tot_tat - prev_tot_tat)) if prev_tot_tat else None
    tot_err_diff = int(cur_tot_err - prev_tot_err) if prev_tot_err is not None else None

    overall_row = {
        "group_name": "OVERALL",
        "supporting_stations": "All Active Stations",
        "volume": cur_tot_u,
        "avg_monthly_output": int(round(cur_tot_u / 3.0)),
        "pi": cur_tot_pi,
        "qoq_pi_str": f"+{tot_pi_diff:.1f}" if (tot_pi_diff and tot_pi_diff > 0) else (f"{tot_pi_diff:.1f}" if (tot_pi_diff and tot_pi_diff < 0) else "(-)"),
        "qoq_pi_val": tot_pi_diff or 0.0,
        "tat": cur_tot_tat,
        "tat_met": cur_tot_tat >= 90,
        "qoq_tat_str": f"+{tot_tat_diff}%" if (tot_tat_diff and tot_tat_diff > 0) else (f"{tot_tat_diff}%" if (tot_tat_diff and tot_tat_diff < 0) else "(-)"),
        "qoq_tat_val": tot_tat_diff or 0,
        "errors": cur_tot_err,
        "qoq_err_str": f"+{tot_err_diff}" if (tot_err_diff and tot_err_diff > 0) else (f"{tot_err_diff}" if (tot_err_diff and tot_err_diff < 0) else "(-)"),
        "qoq_err_val": tot_err_diff or 0,
        "quality": max(0.0, round(((cur_tot_u - cur_tot_err) / cur_tot_u) * 100.0, 1)) if cur_tot_u > 0 else 100.0,
        "quality_met": (cur_tot_u - cur_tot_err) / cur_tot_u >= 0.98 if cur_tot_u > 0 else True
    }

    return quarterly_group_data, overall_row, raw_calibrations_rows, raw_labor_rows, [], []

@app.route("/analytics", strict_slashes=False)
def analytics():
    conn = get_db()
    periods_db = conn.execute("SELECT DISTINCT cal_year as yr, cal_month as mo FROM calibrations WHERE cal_year IS NOT NULL AND cal_month IS NOT NULL ORDER BY yr ASC, mo ASC").fetchall()
    if not periods_db:
        conn.close()
        return render_template("analytics.html", has_data=False)

    timeline, fiscal_quarters_map = [], {}
    for r in periods_db:
        cal_y, m = int(r["yr"]), int(r["mo"])
        q_label, fiscal_yr, full_fy_label = get_keysight_fiscal_info(cal_y, m)
        q_key = f"{q_label} FY'{str(fiscal_yr)[-2:]}"
        if q_key not in fiscal_quarters_map:
            fiscal_quarters_map[q_key] = {"key": q_key, "label": f"{q_key} ({q_label})", "fiscal_yr": fiscal_yr, "q_num": int(q_label.replace("Q", ""))}

        timeline.append({
            "cal_year": cal_y, "fiscal_year": fiscal_yr, "month_num": m,
            "month_name": calendar.month_abbr[m], "label": f"{calendar.month_abbr[m]} '{str(cal_y)[-2:]}",
            "quarter": q_label, "fiscal_year_label": full_fy_label
        })

    available_quarters = list(fiscal_quarters_map.values())
    available_quarters.sort(key=lambda x: (x["fiscal_yr"], x["q_num"]), reverse=True)
    # Inject Full Fiscal Year choices into the dropdown list
    distinct_fy = sorted(list(set(t["fiscal_year"] for t in timeline)), reverse=True)
    for fy in distinct_fy:
        available_quarters.append({
            "key": f"FULL_FY_{fy}",
            "label": f"🌟 FY'{str(fy)[-2:]} Full Year (All Quarters)",
            "fiscal_yr": fy,
            "q_num": 99
        })
    req_quarter = request.args.get("quarter")
    selected_quarter = req_quarter if req_quarter and req_quarter in fiscal_quarters_map else available_quarters[0]["key"]

    active_engineers = conn.execute("SELECT e.id, e.formal_name, e.username, e.station_id, e.shift_hours, e.activation_date, e.resigned_date FROM engineers e WHERE (e.track_performance = 1 OR e.track_performance IS NULL)").fetchall()
    roster_usernames = set([e["username"].strip().lower() for e in active_engineers if e["username"]])
    muted_users = set([m["username"].strip().lower() for m in conn.execute("SELECT username FROM ignored_technicians").fetchall()])

    df_cals_all = pd.read_sql("SELECT category, mos, cal_by, cal_year, cal_month, model_id FROM calibrations", conn)
    if not df_cals_all.empty:
        df_cals_all["clean_by"] = df_cals_all["cal_by"].astype(str).str.strip().str.lower()
        df_cals_all = df_cals_all[df_cals_all["clean_by"].isin(roster_usernames) & (~df_cals_all["clean_by"].isin(muted_users))]

    # Compact 4-Panel Timeline Data
    four_panel_data = {}
    for cat in TAT_CATEGORIES:
        k, lbl = cat["key"], cat["label"]
        vols, mos_pcts = [], []
        for item in timeline:
            cy, m = item["cal_year"], item["month_num"]
            m_rows = df_cals_all[(df_cals_all["category"].astype(str).str.strip().str.upper() == k) & (df_cals_all["cal_year"] == cy) & (df_cals_all["cal_month"] == m)]
            tot_cnt = len(m_rows)
            vols.append(tot_cnt)
            mos_pcts.append(round(((m_rows["mos"] == 1).sum() / tot_cnt) * 100.0, 1) if tot_cnt > 0 else 0.0)

        valid_vols = [v for v in vols if v > 0]
        avg_vol = int(round(sum(valid_vols) / len(valid_vols))) if valid_vols else 0
        valid_mos = [p for p, v in zip(mos_pcts, vols) if v > 0]
        avg_mos = round(sum(valid_mos) / len(valid_mos), 1) if valid_mos else 0.0
        four_panel_data[k] = {"label": lbl, "volumes": vols, "mos_percentages": mos_pcts, "avg_volume": avg_vol, "avg_mos": avg_mos}

    df_att_all = pd.read_sql("SELECT formal_name, capped_hours, year, month FROM attendance", conn)
    if not df_att_all.empty: df_att_all["clean_name"] = df_att_all["formal_name"].apply(normalize_str)
    df_rejs_all = pd.read_sql("SELECT cal_by, reject_year, reject_month FROM rejection_logs", conn)
    if not df_rejs_all.empty:
        df_rejs_all["clean_by"] = df_rejs_all["cal_by"].astype(str).str.strip().str.lower()
        df_rejs_all = df_rejs_all[df_rejs_all["clean_by"].isin(roster_usernames) & (~df_rejs_all["clean_by"].isin(muted_users))]

    monthly_labels, monthly_pi_vols, monthly_pis = [], [], []
    monthly_qualities, monthly_errors = [], []

    for item in timeline:
        cy, cm = item["cal_year"], item["month_num"]
        monthly_labels.append(item["label"])

        c_m = df_cals_all[(df_cals_all["cal_year"] == cy) & (df_cals_all["cal_month"] == cm)]
        u_cnt = len(c_m)
        monthly_pi_vols.append(u_cnt)

        _, m_days = calendar.monthrange(cy, cm)
        m_start, m_end = date(cy, cm, 1), date(cy, cm, m_days)
        tot_days = 0.0
        for e in active_engineers:
            act_d, res_d = parse_iso_date(e["activation_date"]), parse_iso_date(e["resigned_date"])
            if act_d and act_d > m_end: continue
            if res_d and res_d < m_start: continue
            fn_norm = normalize_str(e["formal_name"])
            shift = e["shift_hours"] if e["shift_hours"] > 0 else 8.0
            hrs = df_att_all[(df_att_all["clean_name"] == fn_norm) & (df_att_all["year"] == cy) & (df_att_all["month"] == cm)]["capped_hours"].sum() if not df_att_all.empty else 0.0
            tot_days += (hrs / shift)

        pi_val = round(u_cnt / tot_days, 1) if tot_days > 0 else 0.0
        monthly_pis.append(pi_val)

        r_cnt = len(df_rejs_all[(df_rejs_all["reject_year"] == cy) & (df_rejs_all["reject_month"] == cm)]) if not df_rejs_all.empty else 0
        monthly_errors.append(r_cnt)
        q_pct = max(0.0, round(((u_cnt - r_cnt) / u_cnt) * 100.0, 1)) if u_cnt > 0 else 100.0
        monthly_qualities.append(q_pct)

    quarterly_group_data, overall_row, _, _, _, unmapped_models_in_quarter = calculate_quarterly_engine(conn, selected_quarter)
    target_groups = ["CSG", "EiSG+Scope", "STDS LAB", "Legacy+Modular+PMPS (A)+MMA+ Others", "MEC Lab"]
    conn.close()

    return render_template(
        "analytics.html", has_data=True, timeline_json=json.dumps(timeline), four_panel_data=four_panel_data,
        available_quarters=available_quarters, selected_quarter=selected_quarter,
        quarterly_group_data=quarterly_group_data, overall_row=overall_row,
        unmapped_models_in_quarter=unmapped_models_in_quarter, target_groups=target_groups,
        monthly_labels=json.dumps(monthly_labels), monthly_pi_vols=json.dumps(monthly_pi_vols),
        monthly_pis=json.dumps(monthly_pis), monthly_qualities=json.dumps(monthly_qualities),
        monthly_errors=json.dumps(monthly_errors)
    )

@app.route("/export_quarterly_data", strict_slashes=False)
def export_quarterly_data():
    conn = get_db()
    req_quarter = request.args.get("quarter")
    periods_db = conn.execute("SELECT DISTINCT cal_year as yr, cal_month as mo FROM calibrations WHERE cal_year IS NOT NULL AND cal_month IS NOT NULL ORDER BY yr ASC, mo ASC").fetchall()
    if not periods_db:
        conn.close()
        flash("No calibration data available.", "warning")
        return redirect(url_for("analytics"))

    fiscal_quarters_map = {}
    for r in periods_db:
        cal_y, m = int(r["yr"]), int(r["mo"])
        q_label, fiscal_yr, _ = get_keysight_fiscal_info(cal_y, m)
        q_key = f"{q_label} FY'{str(fiscal_yr)[-2:]}"
        fiscal_quarters_map[q_key] = True
        fiscal_quarters_map[f"FULL_FY_{fiscal_yr}"] = True

    selected_quarter = req_quarter if req_quarter and req_quarter in fiscal_quarters_map else list(fiscal_quarters_map.keys())[-1]
    quarterly_group_data, overall_row, raw_calibrations_rows, raw_labor_rows, _, _ = calculate_quarterly_engine(conn, selected_quarter)
    conn.close()

    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    sub_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    overall_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    white_font, reg_font = Font(name="Arial", size=10, bold=True, color="FFFFFF"), Font(name="Arial", size=9)
    bold_dark = Font(name="Arial", size=9, bold=True, color="000000")
    center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin', color='B0C4DE'), right=Side(style='thin', color='B0C4DE'), top=Side(style='thin', color='B0C4DE'), bottom=Side(style='thin', color='B0C4DE'))

    ws_sum = wb.active
    ws_sum.title = "Quarterly Summary"
    ws_sum.append(["Segment Group", "Assigned Stations", "Total Volume", "Avg Monthly Output", "Quarterly P.I.", "QoQ Δ P.I.", "Weighted TAT %", "QoQ Δ TAT", "Rejections", "QoQ Δ Rej", "Quality %"])
    for c_i in range(1, 12):
        c = ws_sum.cell(row=1, column=c_i)
        c.fill, c.font, c.alignment = header_fill, white_font, center

    r_idx = 2
    for g in quarterly_group_data:
        ws_sum.append([
            g["group_name"], g["supporting_stations"], g["volume"], 
            g["avg_monthly_output"], g["pi"], g["qoq_pi_str"], 
            f"{g['tat']}%", g["qoq_tat_str"], g["errors"], 
            g["qoq_err_str"], f"{g['quality']}%"
        ])
        for c_i in range(1, 12):
            cell = ws_sum.cell(row=r_idx, column=c_i)
            cell.font, cell.alignment, cell.border = reg_font, center, thin_border
        r_idx += 1

    ws_sum.append([
        overall_row["group_name"], overall_row["supporting_stations"], overall_row["volume"],
        overall_row["avg_monthly_output"], overall_row["pi"], overall_row["qoq_pi_str"],
        f"{overall_row['tat']}%", overall_row["qoq_tat_str"], overall_row["errors"],
        overall_row["qoq_err_str"], f"{overall_row['quality']}%"
    ])
    for c_i in range(1, 12):
        cell = ws_sum.cell(row=r_idx, column=c_i)
        cell.fill, cell.font, cell.alignment, cell.border = overall_fill, bold_dark, center, thin_border

    ws_cal = wb.create_sheet(title="Calibrations Breakdown")
    ws_cal.append(["Service ID", "Model ID", "Serial No", "Cal Date", "Month", "Year", "CalBy", "Category", "MOS", "Assigned Group", "AE Station"])
    for c_i in range(1, 12):
        c = ws_cal.cell(row=1, column=c_i)
        c.fill, c.font, c.alignment = sub_fill, white_font, center
    for r in raw_calibrations_rows: ws_cal.append(r)

    ws_labor = wb.create_sheet(title="Labor & Working Days")
    ws_labor.append(["Engineer Name", "Username", "Station", "Year", "Month", "Working Hours", "Shift", "Effective Days", "Assigned Group"])
    for c_i in range(1, 10):
        c = ws_labor.cell(row=1, column=c_i)
        c.fill, c.font, c.alignment = sub_fill, white_font, center
    for r in raw_labor_rows: ws_labor.append(r)

    for sheet in [ws_sum, ws_cal, ws_labor]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 3, 14)

    conn.close()
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    clean_q = re.sub(r'[^a-zA-Z0-9_]', '_', selected_quarter)
    return send_file(excel_stream, as_attachment=True, download_name=f"Quarterly_Calculation_Audit_{clean_q}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/export_pptx", strict_slashes=False)
def export_pptx():
    conn = get_db()
    req_period = request.args.get("period", "2026-08")
    try: y, m = [int(x) for x in req_period.split("-")]
    except ValueError: y, m = 2026, 8
    ppt_stream = generate_monthly_deck(conn, y, m)
    conn.close()
    return send_file(ppt_stream, as_attachment=True, download_name=f"Cal_Lab_Output_Tracking_{calendar.month_abbr[m]}_{y}.pptx", mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")

@app.route("/history", strict_slashes=False)
def history():
    conn = get_db()
    mode = request.args.get("mode", "engineer")
    show_resigned = request.args.get("show_resigned") == "1"
    today = date.today()
    
    all_engineers = conn.execute("SELECT e.id, e.formal_name, e.display_name, e.username, e.station_id, e.custom_target_pi, s.name as station_name, e.shift_hours, e.joined_date, e.activation_date, e.resigned_date FROM engineers e LEFT JOIN stations s ON e.station_id = s.id WHERE (e.track_performance = 1 OR e.track_performance IS NULL) ORDER BY e.formal_name").fetchall()
    filtered_engineers = []
    for e in all_engineers:
        res_d = parse_iso_date(e["resigned_date"])
        is_resigned = res_d and res_d <= today
        e_dict = dict(e)
        e_dict["is_resigned"] = is_resigned
        if not is_resigned or show_resigned: filtered_engineers.append(e_dict)

    stations = conn.execute("SELECT id, name, sort_order, target_pi FROM stations WHERE is_active = 1 ORDER BY sort_order ASC, name ASC").fetchall()
    valid_station_names = [s["name"] for s in stations]
    periods = conn.execute("SELECT DISTINCT cal_year as yr, cal_month as mo FROM calibrations WHERE cal_year IS NOT NULL UNION SELECT DISTINCT year as yr, month as mo FROM attendance WHERE year IS NOT NULL ORDER BY yr ASC, mo ASC").fetchall()

    df_cals_all = pd.read_sql("SELECT cal_by, category, mos, cal_year, cal_month FROM calibrations", conn)
    df_att_all = pd.read_sql("SELECT formal_name, year, month, capped_hours FROM attendance", conn)
    if not df_att_all.empty: df_att_all["clean_name"] = df_att_all["formal_name"].apply(normalize_str)
    df_mistakes_all = pd.read_sql("SELECT engineer_id, year, month, mistakes FROM quality_logs", conn)

    if mode == "engineer":
        target_id = request.args.get("target")
        selected_eng = next((e for e in filtered_engineers if str(e["id"]) == str(target_id)), None) if target_id else None
        if not selected_eng and filtered_engineers: selected_eng = filtered_engineers[0]

        selected_target = str(selected_eng["id"]) if selected_eng else ""
        shift_base = selected_eng["shift_hours"] if selected_eng else 8.0
        entity_list = filtered_engineers

        hist_records = conn.execute("SELECT h.effective_year, h.effective_month, s.name as station_name FROM engineer_station_history h JOIN stations s ON h.station_id = s.id WHERE h.engineer_id = ? ORDER BY h.effective_year ASC, h.effective_month ASC", (selected_eng["id"],)).fetchall() if selected_eng else []
        transfer_timeline = []

        if hist_records:
            first_h = hist_records[0]
            first_mo = calendar.month_abbr[first_h['effective_month']]
            first_yr = str(first_h['effective_year'])[-2:]
            transfer_timeline.append(f"{first_h['station_name']} ({first_mo} '{first_yr})")
            last_station = first_h['station_name']
            for hr in hist_records[1:]:
                curr_station = hr['station_name']
                if curr_station != last_station:
                    mo_str = calendar.month_abbr[hr['effective_month']]
                    yr_str = str(hr['effective_year'])[-2:]
                    transfer_timeline.append(f"➔ {curr_station} ({mo_str} '{yr_str})")
                    last_station = curr_station
        elif selected_eng:
            transfer_timeline.append(f"{selected_eng['station_name'] or 'Unassigned'} (Initial)")

        ledger, chart_labels, chart_units, chart_pis, chart_qualities, chart_tats = [], [], [], [], [], []
        chart_ete, chart_new, chart_psc_ext, chart_psc_int = [], [], [], []
        tot_u, tot_m, pi_sum, q_sum, tat_sum, valid_cnt = 0, 0, 0.0, 0.0, 0.0, 0
        latest_station_name = selected_eng["station_name"] if selected_eng else "Unassigned"
        base_target_pi = 3.0
        distinct_years = sorted(list(set([int(p["yr"]) for p in periods])), reverse=True)

        if selected_eng:
            prev_st_name = None
            for p in periods:
                yr, mo = int(p["yr"]), int(p["mo"])
                p_label = f"{calendar.month_abbr[mo]} '{str(yr)[-2:]}"
                _, st_name, _, st_target, _ = resolve_engineer_station(conn, selected_eng["id"], selected_eng["station_id"], yr, mo)
                latest_station_name = st_name
                ind_target = selected_eng["custom_target_pi"] if selected_eng["custom_target_pi"] is not None else st_target
                base_target_pi = ind_target
                is_transfer_month = (prev_st_name is not None and prev_st_name != st_name)
                transfer_note = f"Transferred from {prev_st_name}" if is_transfer_month else None
                prev_st_name = st_name

                p_cals = df_cals_all[(df_cals_all["cal_by"] == selected_eng["username"]) & (df_cals_all["cal_year"] == yr) & (df_cals_all["cal_month"] == mo)] if not df_cals_all.empty else pd.DataFrame()
                units = len(p_cals)
                hrs = df_att_all[(df_att_all["clean_name"] == normalize_str(selected_eng["formal_name"])) & (df_att_all["year"] == yr) & (df_att_all["month"] == mo)]["capped_hours"].sum() if not df_att_all.empty else 0.0
                equiv_days = int(round(hrs / shift_base)) if shift_base > 0 else 0
                pi = round(units / (hrs / shift_base), 1) if hrs > 0 else 0.0

                m_rows = df_mistakes_all[(df_mistakes_all["engineer_id"] == selected_eng["id"]) & (df_mistakes_all["year"] == yr) & (df_mistakes_all["month"] == mo)]
                mistakes = int(m_rows["mistakes"].values[0]) if not m_rows.empty else 0
                quality = max(0.0, round(((units - mistakes) / units) * 100.0, 1)) if units > 0 else 100.0

                cat_scores, cat_numeric = {}, {}
                w_num, act_w = 0.0, 0.0
                for cat in TAT_CATEGORIES:
                    k, w, lbl = cat["key"], cat["weight"], cat["label"].lower()
                    c_rows = p_cals[p_cals["category"].str.strip().str.upper() == k] if not p_cals.empty else pd.DataFrame()
                    if len(c_rows) > 0:
                        pct = int(round(((c_rows["mos"] == 1).sum() / len(c_rows)) * 100.0))
                        cat_scores[lbl], cat_numeric[lbl] = f"{pct}%", pct
                        w_num += pct * (w / 100.0); act_w += (w / 100.0)
                    else: cat_scores[lbl], cat_numeric[lbl] = "N/A", None

                tat_val = int(round(w_num / act_w)) if act_w > 0 else None
                tat_str = f"{tat_val}%" if tat_val is not None else "N/A"

                if units > 0 or hrs > 0:
                    ledger.append({
                        "year": yr, "month_num": mo, "period_label": f"{calendar.month_name[mo]} {yr}",
                        "station": st_name, "transfer_note": transfer_note, "target_pi": ind_target,
                        "units": units, "hours": round(hrs, 1), "equiv_days": equiv_days,
                        "pi": pi, "pi_met": pi >= ind_target, "mistakes": mistakes, "quality": quality,
                        "ete": cat_scores.get("ete", "N/A"), "ete_num": cat_numeric.get("ete"),
                        "new": cat_scores.get("new", "N/A"), "new_num": cat_numeric.get("new"),
                        "psc_ext": cat_scores.get("psc_ext", "N/A"), "psc_ext_num": cat_numeric.get("psc_ext"),
                        "psc_int": cat_scores.get("psc_int", "N/A"), "psc_int_num": cat_numeric.get("psc_int"),
                        "tat": tat_str, "tat_num": tat_val
                    })
                    chart_labels.append(p_label); chart_units.append(units); chart_pis.append(pi); chart_qualities.append(quality); chart_tats.append(tat_val or 0)
                    chart_ete.append(cat_numeric.get("ete", 0)); chart_new.append(cat_numeric.get("new", 0)); chart_psc_ext.append(cat_numeric.get("psc_ext", 0)); chart_psc_int.append(cat_numeric.get("psc_int", 0))
                    tot_u += units; tot_m += mistakes; pi_sum += pi; q_sum += quality
                    if tat_val is not None: tat_sum += tat_val
                    valid_cnt += 1

            profile = {
                "id": selected_eng["id"], "title": selected_eng["formal_name"], "subtitle": f"{latest_station_name} Station",
                "timeline": " ".join(transfer_timeline) if len(transfer_timeline) > 1 else None,
                "shift": f"{selected_eng['shift_hours']}h", "joined": format_display_date(selected_eng["joined_date"]),
                "activated": format_display_date(selected_eng["activation_date"]), "resigned": format_display_date(selected_eng["resigned_date"]) if selected_eng["resigned_date"] else None,
                "tenure": calculate_tenure(selected_eng["joined_date"] or selected_eng["activation_date"], selected_eng["resigned_date"]),
                "base_target_pi": base_target_pi
            }
        else:
            profile = {"id": None, "title": "No Engineer Selected", "subtitle": "Unassigned", "timeline": None, "shift": "8h", "joined": "N/A", "activated": "N/A", "resigned": None, "tenure": "N/A", "base_target_pi": 3.0}

        avg_pi = round(pi_sum / valid_cnt, 1) if valid_cnt > 0 else 0.0
        avg_quality = round(q_sum / valid_cnt, 1) if valid_cnt > 0 else 100.0
        avg_tat = round(tat_sum / valid_cnt, 1) if valid_cnt > 0 else 0.0
        active_ledger = [r for r in ledger if r["units"] > 0 or r["hours"] > 0]
        pi_met_cnt = sum(1 for r in active_ledger if r["pi_met"])
        pi_hit_rate = round((pi_met_cnt / len(active_ledger)) * 100.0, 1) if active_ledger else 0.0

        cat_averages = {}
        for cat_k in ["ete_num", "new_num", "psc_ext_num", "psc_int_num"]:
            valid_sc = [r[cat_k] for r in active_ledger if r[cat_k] is not None]
            cat_averages[cat_k.replace("_num", "").upper()] = round(sum(valid_sc) / len(valid_sc), 1) if valid_sc else 0.0

        top_cat = max(cat_averages.items(), key=lambda x: x[1]) if cat_averages else ("N/A", 0)
        low_cat = min(cat_averages.items(), key=lambda x: x[1]) if cat_averages else ("N/A", 0)

        appraisal_summary = {
            "pi_hit_rate": pi_hit_rate, "pi_met_cnt": pi_met_cnt, "total_months": len(active_ledger),
            "defect_rate": round((tot_m / tot_u * 100.0), 2) if tot_u > 0 else 0.0,
            "top_category": top_cat[0], "top_category_rate": top_cat[1], "low_category": low_cat[0], "low_category_rate": low_cat[1]
        }

        lifetime = {"total_units": tot_u, "total_mistakes": tot_m, "avg_pi": avg_pi, "avg_quality": avg_quality, "avg_tat": avg_tat, "pi_diff": round(avg_pi - base_target_pi, 1), "tat_diff": round(avg_tat - 90.0, 1)}
        station_peers_ledger, station_ae_output_datasets, station_ae_pi_datasets = [], [], []
        ledger = list(reversed(ledger))

    else:
        target_name = request.args.get("target")
        selected_st = target_name if (target_name and target_name in valid_station_names) else (valid_station_names[0] if valid_station_names else "Signal Analyzer")
        st_obj = next((s for s in stations if s["name"] == selected_st), None)
        st_target = st_obj["target_pi"] if st_obj else 3.0
        selected_target, entity_list, distinct_years = selected_st, stations, []
        profile = {"id": None, "title": f"{selected_st} Station", "subtitle": f"Station Default Target P.I.: {st_target} units/day", "timeline": None, "shift": None, "joined": None, "activated": None, "resigned": None, "tenure": None, "base_target_pi": st_target}
        appraisal_summary = {}

        st_peers_all = [e for e in all_engineers if e["station_name"] == selected_st or conn.execute("SELECT 1 FROM engineer_station_history WHERE engineer_id = ? AND station_id = (SELECT id FROM stations WHERE name = ?)", (e["id"], selected_st)).fetchone()]
        st_colors = ['#0d6efd', '#20c997', '#fd7e14', '#6f42c1', '#d63384', '#0dcaf0', '#198754']
        chart_labels = [f"{calendar.month_abbr[int(p['mo'])]} '{str(p['yr'])[-2:]}" for p in periods]
        station_ae_output_datasets, station_ae_pi_datasets = [], []
        tot_u, tot_m, pi_sum, valid_cnt = 0, 0, 0.0, 0

        for idx, e in enumerate(st_peers_all):
            ae_legend_name = e["display_name"] or e["formal_name"].split()[0]
            ae_color = st_colors[idx % len(st_colors)]
            ae_units_series, ae_pi_series = [], []
            act_d, res_d = parse_iso_date(e["activation_date"]), parse_iso_date(e["resigned_date"])

            for p in periods:
                yr, mo = int(p["yr"]), int(p["mo"])
                p_start, p_end = date(yr, mo, 1), date(yr, mo, calendar.monthrange(yr, mo)[1])
                _, st_name, _, _, _ = resolve_engineer_station(conn, e["id"], e["station_id"], yr, mo)

                if (act_d and act_d > p_end) or (res_d and res_d < p_start) or (st_name != selected_st):
                    ae_units_series.append(None); ae_pi_series.append(None)
                else:
                    p_cals = df_cals_all[(df_cals_all["cal_by"] == e["username"]) & (df_cals_all["cal_year"] == yr) & (df_cals_all["cal_month"] == mo)] if not df_cals_all.empty else pd.DataFrame()
                    units = len(p_cals)
                    hrs = df_att_all[(df_att_all["clean_name"] == normalize_str(e["formal_name"])) & (df_att_all["year"] == yr) & (df_att_all["month"] == mo)]["capped_hours"].sum() if not df_att_all.empty else 0.0
                    pi = round(units / (hrs / e["shift_hours"]), 1) if (hrs > 0 and e["shift_hours"] > 0) else 0.0
                    ae_units_series.append(units); ae_pi_series.append(pi)
                    tot_u += units
                    if pi > 0: pi_sum += pi; valid_cnt += 1

            station_ae_output_datasets.append({"label": ae_legend_name, "data": ae_units_series, "backgroundColor": ae_color, "borderRadius": 3})
            station_ae_pi_datasets.append({"label": ae_legend_name, "data": ae_pi_series, "borderColor": ae_color, "backgroundColor": ae_color, "borderWidth": 2, "pointRadius": 3, "tension": 0.2})

        station_ae_pi_datasets.append({"label": f"Station Target ({st_target})", "data": [st_target] * len(periods), "borderColor": "#dc3545", "borderWidth": 2, "borderDash": [5, 5], "pointRadius": 0, "fill": False})

        station_peers_ledger = []
        for p in reversed(periods):
            yr, mo = int(p["yr"]), int(p["mo"])
            month_peers = []
            for e in st_peers_all:
                _, st_name, _, _, _ = resolve_engineer_station(conn, e["id"], e["station_id"], yr, mo)
                if st_name == selected_st:
                    ind_target = e["custom_target_pi"] if e["custom_target_pi"] is not None else st_target
                    u, fn, shift = e["username"], normalize_str(e["formal_name"]), e["shift_hours"]
                    u_cals = df_cals_all[(df_cals_all["cal_by"] == u) & (df_cals_all["cal_year"] == yr) & (df_cals_all["cal_month"] == mo)] if not df_cals_all.empty else pd.DataFrame()
                    units = len(u_cals)
                    hrs = df_att_all[(df_att_all["clean_name"] == fn) & (df_att_all["year"] == yr) & (df_att_all["month"] == mo)]["capped_hours"].sum() if not df_att_all.empty else 0.0
                    equiv_days = int(round(hrs / shift)) if shift > 0 else 0
                    pi = round(units / (hrs / shift), 1) if hrs > 0 else 0.0

                    m_rows = df_mistakes_all[(df_mistakes_all["engineer_id"] == e["id"]) & (df_mistakes_all["year"] == yr) & (df_mistakes_all["month"] == mo)]
                    mistakes = int(m_rows["mistakes"].values[0]) if not m_rows.empty else 0
                    quality = max(0.0, round(((units - mistakes) / units) * 100.0, 1)) if units > 0 else 100.0

                    w_num, act_w = 0.0, 0.0
                    for cat in TAT_CATEGORIES:
                        k, w = cat["key"], cat["weight"]
                        c_rows = u_cals[u_cals["category"].str.strip().str.upper() == k] if not u_cals.empty else pd.DataFrame()
                        if len(c_rows) > 0:
                            pct = int(round(((c_rows["mos"] == 1).sum() / len(c_rows)) * 100.0))
                            w_num += pct * (w / 100.0); act_w += (w / 100.0)
                    tat_val = int(round(w_num / act_w)) if act_w > 0 else None

                    if units > 0 or hrs > 0:
                        month_peers.append({"name": e["formal_name"], "units": units, "hours": round(hrs, 1), "equiv_days": equiv_days, "target_pi": ind_target, "pi": pi, "pi_met": pi >= ind_target, "mistakes": mistakes, "quality": quality, "tat": f"{tat_val}%" if tat_val is not None else "N/A"})
                        tot_m += mistakes

            if month_peers:
                st_month_units = sum(p["units"] for p in month_peers)
                st_month_days = sum(p["equiv_days"] for p in month_peers)
                st_month_pi = round(st_month_units / st_month_days, 1) if st_month_days > 0 else 0.0
                station_peers_ledger.append({"period_label": f"{calendar.month_name[mo]} {yr}", "station_units": st_month_units, "station_pi": st_month_pi, "station_pi_met": st_month_pi >= st_target, "peers": month_peers})

        lifetime = {"total_units": tot_u, "total_mistakes": tot_m, "avg_pi": round(pi_sum / valid_cnt, 1) if valid_cnt > 0 else 0.0, "avg_quality": 99.0, "avg_tat": 92.0, "pi_diff": 0.0, "tat_diff": 0.0}
        ledger, chart_units, chart_pis, chart_qualities, chart_tats = [], [], [], [], []
        chart_ete, chart_new, chart_psc_ext, chart_psc_int = [], [], [], []

    conn.close()
    return render_template(
        "history.html", mode=mode, selected_target=selected_target, show_resigned=show_resigned,
        entity_list=entity_list, profile=profile, lifetime=lifetime, appraisal_summary=appraisal_summary,
        ledger=ledger, distinct_years=distinct_years, station_peers_ledger=station_peers_ledger,
        station_ae_output_datasets=json.dumps(station_ae_output_datasets),
        station_ae_pi_datasets=json.dumps(station_ae_pi_datasets),
        chart_labels=json.dumps(chart_labels), chart_units=json.dumps(chart_units), chart_pis=json.dumps(chart_pis),
        chart_qualities=json.dumps(chart_qualities), chart_tats=json.dumps(chart_tats),
        chart_ete=json.dumps(chart_ete), chart_new=json.dumps(chart_new), chart_psc_ext=json.dumps(chart_psc_ext), chart_psc_int=json.dumps(chart_psc_int)
    )

@app.route("/upload", strict_slashes=False)
def upload():
    return render_template("upload.html")

@app.route("/upload/inspect", methods=["POST"])
def upload_inspect():
    conn = get_db()
    cal_file, att_file, rej_file = request.files.get("calibration_file"), request.files.get("attendance_file"), request.files.get("rejection_file")
    has_cal, has_att, has_rej = cal_file and cal_file.filename != "", att_file and att_file.filename != "", rej_file and rej_file.filename != ""

    if not has_cal and not has_att and not has_rej:
        flash("Please select at least one file to inspect.", "warning")
        return redirect(url_for("upload"))

    cal_summary, att_summary, rej_summary, unmapped_models = None, None, None, []
    existing_groups = [r[0] for r in conn.execute("SELECT DISTINCT group_name FROM model_groups ORDER BY group_name").fetchall()] or ["CSG", "EiSG+Scope", "STDS LAB", "Legacy+Modular+PMPS (A)+MMA+ Others", "MEC Lab"]

    all_engineers = conn.execute("SELECT * FROM engineers WHERE is_active = 1").fetchall()
    roster_usernames = set([e["username"].strip().lower() for e in all_engineers])
    roster_names = set([normalize_str(e["formal_name"]) for e in all_engineers])
    muted_usernames = set([m["username"].strip().lower() for m in conn.execute("SELECT username FROM ignored_technicians").fetchall()])

    if has_cal:
        cal_path = os.path.join(TEMP_UPLOAD_DIR, "pending_cal.xlsx")
        cal_file.save(cal_path)
        df_cal = pd.read_excel(cal_path)
        cols = {str(c).strip().lower(): c for c in df_cal.columns}
        cal_date_col, cal_by_col = cols.get("caldate", "CalDate"), cols.get("calby", "CalBy")
        model_col = cols.get("modelid", cols.get("model_id", cols.get("model id", "ModelID")))

        df_cal = df_cal.dropna(subset=[cal_by_col, cal_date_col])
        df_cal[cal_date_col] = pd.to_datetime(df_cal[cal_date_col], errors="coerce")
        df_cal = df_cal.dropna(subset=[cal_date_col])
        latest_date = df_cal[cal_date_col].max()
        period_label = f"{calendar.month_name[latest_date.month]} {latest_date.year}"

        matched_units, muted_units, muted_count, unmapped_units = 0, 0, 0, 0
        matched_users, unmapped_techs = set(), []
        for user, cnt in df_cal[cal_by_col].astype(str).str.strip().value_counts().to_dict().items():
            u_clean = user.lower()
            if u_clean in roster_usernames: matched_units += cnt; matched_users.add(u_clean)
            elif u_clean in muted_usernames: muted_units += cnt; muted_count += 1
            else: unmapped_units += cnt; unmapped_techs.append({"username": user, "units": cnt})

        mapped_db_models = set([r[0].upper() for r in conn.execute("SELECT UPPER(model_id) FROM model_groups").fetchall()])
        file_models = df_cal[model_col].dropna().astype(str).str.strip().str.upper().unique().tolist() if model_col in df_cal.columns else []
        for m in sorted(file_models):
            if m and m not in mapped_db_models and m != "NAN": unmapped_models.append(m)

        cal_summary = {
            "filename": cal_file.filename, "period_label": period_label, "total_rows": len(df_cal),
            "matched_units": matched_units, "matched_ae_count": len(matched_users), "muted_units": muted_units,
            "muted_count": muted_count, "unmapped_units": unmapped_units, "unmapped_techs": unmapped_techs,
            "unmapped_models_count": len(unmapped_models)
        }

    if has_att:
        att_path = os.path.join(TEMP_UPLOAD_DIR, "pending_att.xlsx")
        att_file.save(att_path)
        with open(att_path, "rb") as f:
            header_idx, year, month = find_attendance_header_and_period(f)
        df_att = pd.read_excel(att_path, skiprows=header_idx)
        df_att.columns = [str(c).strip() for c in df_att.columns]
        file_staff_names = set([normalize_str(str(r.get("Name", "")).strip().replace('"', '')) for _, r in df_att.iterrows() if str(r.get("Name", "")).strip().lower() not in ["nan", "none", ""]])

        matched_roster = roster_names.intersection(file_staff_names)
        missing_roster_names = [e["formal_name"] for e in all_engineers if normalize_str(e["formal_name"]) in roster_names.difference(file_staff_names)]
        att_summary = {"filename": att_file.filename, "period_label": f"{calendar.month_name[month]} {year}", "total_staff": len(file_staff_names), "matched_roster_count": len(matched_roster), "active_roster_total": len(roster_names), "unmatched_staff_count": len(file_staff_names) - len(matched_roster), "missing_roster_names": missing_roster_names}

    if has_rej:
        rej_path = os.path.join(TEMP_UPLOAD_DIR, "pending_rej.xlsx")
        rej_file.save(rej_path)
        df_rej = pd.read_excel(rej_path)
        cols = {str(c).strip().lower(): c for c in df_rej.columns}
        date_col, cal_by_col = cols.get("approved/reject date", cols.get("approved/rejected date", cols.get("caldate", "Approved/Reject Date"))), cols.get("calby", "CalBy")
        df_rej[date_col] = pd.to_datetime(df_rej[date_col], errors="coerce")
        df_rej = df_rej.dropna(subset=[date_col])
        latest_date = df_rej[date_col].max()
        rej_summary = {"filename": rej_file.filename, "period_label": f"{calendar.month_name[latest_date.month]} {latest_date.year}" if not df_rej.empty else "N/A", "total_errors": len(df_rej), "affected_users": len(df_rej[cal_by_col].unique()) if not df_rej.empty else 0}

    conn.close()
    return render_template(
        "preflight.html", cal_summary=cal_summary, att_summary=att_summary, 
        rej_summary=rej_summary, unmapped_models=unmapped_models, existing_groups=existing_groups
    )

@app.route("/upload/confirm", methods=["POST"])
def upload_confirm():
    conn = get_db()
    cur = conn.cursor()
    assigned_models = request.form.getlist("new_model_id")
    for m in assigned_models:
        group_val = request.form.get(f"group_for_{m}", "").strip()
        custom_group = request.form.get(f"custom_group_for_{m}", "").strip()
        final_group = custom_group if group_val == "__NEW__" and custom_group else group_val
        if m and final_group:
            cur.execute("INSERT INTO model_groups (model_id, group_name) VALUES (?, ?) ON CONFLICT(model_id) DO UPDATE SET group_name = excluded.group_name", (m.strip().upper(), final_group.strip()))
    conn.commit()
    conn.close()

    cal_path, att_path, rej_path = os.path.join(TEMP_UPLOAD_DIR, "pending_cal.xlsx"), os.path.join(TEMP_UPLOAD_DIR, "pending_att.xlsx"), os.path.join(TEMP_UPLOAD_DIR, "pending_rej.xlsx")
    committed_items = []
    if os.path.exists(cal_path):
        with open(cal_path, "rb") as f: ingest_calibrations(f)
        os.remove(cal_path); committed_items.append("Calibrations")
    if os.path.exists(att_path):
        with open(att_path, "rb") as f: ingest_attendance(f)
        os.remove(att_path); committed_items.append("Attendance")
    if os.path.exists(rej_path):
        with open(rej_path, "rb") as f: ingest_rejections(f)
        os.remove(rej_path); committed_items.append("Outgoing Inspection Errors")

    flash(f"✅ Success! {' and '.join(committed_items)} committed to database.", "success")
    return redirect(url_for("dashboard"))

@app.route("/upload/cancel", strict_slashes=False)
def upload_cancel():
    for p in [os.path.join(TEMP_UPLOAD_DIR, f) for f in ["pending_cal.xlsx", "pending_att.xlsx", "pending_rej.xlsx"]]:
        if os.path.exists(p): os.remove(p)
    flash("Upload inspection cancelled.", "info")
    return redirect(url_for("upload"))

@app.route("/map_model_group", methods=["POST"])
def map_model_group():
    conn = get_db()
    cur = conn.cursor()
    quarter = request.form.get("quarter", "")
    models = request.form.getlist("model_id")
    groups = request.form.getlist("group_name")

    saved_count = 0
    for m, g in zip(models, groups):
        clean_m, clean_g = m.strip().upper(), g.strip()
        if clean_m and clean_g:
            cur.execute("INSERT INTO model_groups (model_id, group_name) VALUES (?, ?) ON CONFLICT(model_id) DO UPDATE SET group_name = excluded.group_name", (clean_m, clean_g))
            saved_count += 1

    conn.commit()
    conn.close()
    flash(f"Successfully mapped {saved_count} model(s) permanently!", "success")
    return redirect(url_for("analytics", quarter=quarter))

# =============================================================================
# INDIVIDUAL STAFF EDIT & TRANSFER HANDLERS
# =============================================================================
@app.route("/team/edit_member", methods=["POST"], strict_slashes=False)
def edit_team_member():
    conn = get_db()
    eng_id = request.form.get("eng_id")
    formal_name = request.form.get("formal_name", "").strip()
    display_name = request.form.get("display_name", "").strip() or None
    username = request.form.get("username", "").strip()
    shift_hours = float(request.form.get("shift_hours", 8.0))
    custom_target = request.form.get("custom_target_pi", "").strip()
    custom_target_pi = float(custom_target) if custom_target else None
    track_perf = 1 if request.form.get("track_performance") else 0
    joined_date = request.form.get("joined_date", "").strip() or None
    activation_date = request.form.get("activation_date", "").strip() or None
    resigned_date = request.form.get("resigned_date", "").strip() or None

    if not username:
        clean_name = re.sub(r'[^a-z0-9]', '.', formal_name.lower()).strip('.')
        username = f"admin.{clean_name}"

    conn.execute("""
        UPDATE engineers 
        SET formal_name = ?, display_name = ?, username = ?, shift_hours = ?, 
            custom_target_pi = ?, track_performance = ?, joined_date = ?, 
            activation_date = ?, resigned_date = ?
        WHERE id = ?
    """, (formal_name, display_name, username, shift_hours, custom_target_pi, track_perf, joined_date, activation_date, resigned_date, eng_id))
    conn.commit()
    conn.close()
    flash(f"Profile for '{formal_name}' updated successfully!", "success")
    return redirect(url_for("team"))


@app.route("/team/transfer_member", methods=["POST"], strict_slashes=False)
def transfer_team_member():
    conn = get_db()
    eng_id = request.form.get("eng_id")
    new_station_id = request.form.get("new_station_id")
    transfer_period = request.form.get("transfer_period", "").strip()

    if not eng_id or not new_station_id or not transfer_period:
        flash("Transfer requires a station and an effective month.", "warning")
        return redirect(url_for("team"))

    new_st_id = int(new_station_id)
    cur_eng = conn.execute("SELECT formal_name, station_id, joined_date, activation_date FROM engineers WHERE id = ?", (eng_id,)).fetchone()
    if not cur_eng:
        conn.close()
        return redirect(url_for("team"))

    old_station_id = cur_eng["station_id"]
    formal_name = cur_eng["formal_name"]

    match = re.match(r'(\d{4})-(\d{2})', transfer_period)
    if match:
        t_yr, t_mo = int(match.group(1)), int(match.group(2))
        start_d = parse_iso_date(cur_eng["activation_date"] or cur_eng["joined_date"]) or date(2024, 1, 1)

        # 1. Lock in Initial Station starting from their join date
        if old_station_id:
            conn.execute("""
                INSERT OR IGNORE INTO engineer_station_history (engineer_id, station_id, effective_year, effective_month)
                VALUES (?, ?, ?, ?)
            """, (eng_id, old_station_id, start_d.year, start_d.month))

        # 2. Record New Station for the transfer month
        conn.execute("""
            INSERT INTO engineer_station_history (engineer_id, station_id, effective_year, effective_month)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(engineer_id, effective_year, effective_month)
            DO UPDATE SET station_id = excluded.station_id
        """, (eng_id, new_st_id, t_yr, t_mo))

        # 3. Update main profile
        conn.execute("UPDATE engineers SET station_id = ? WHERE id = ?", (new_st_id, eng_id))
        conn.commit()

        new_st_name = conn.execute("SELECT name FROM stations WHERE id = ?", (new_st_id,)).fetchone()[0]
        flash(f"Successfully transferred '{formal_name}' to {new_st_name} effective {transfer_period}!", "success")

    conn.close()
    return redirect(url_for("team"))

@app.route("/team", methods=["GET", "POST"], strict_slashes=False)
def team():
    conn = get_db()
    if request.method == "POST":
        action = request.form.get("action")

        if action == "update_kpis":
            target_label = request.form.get("target_label", "Q2 FY'26 Target").strip()
            target_hc = request.form.get("target_headcount", "23").strip()
            conn.execute("INSERT OR REPLACE INTO system_settings (setting_key, setting_val) VALUES ('target_headcount', ?)", (target_hc,))
            for k_key in request.form.getlist("kpi_key"):
                val = float(request.form.get(f"val_{k_key}", 0.0))
                conn.execute("UPDATE kpi_configs SET target_label = ?, target_value = ? WHERE kpi_key = ?", (target_label, val, k_key))
            conn.commit()
            flash("KPI Targets and Target Headcount updated successfully!", "success")

        elif action == "update_stations":
            for st_id in request.form.getlist("station_id"):
                st_name = request.form.get(f"name_{st_id}", "").strip().upper()
                grp_name = request.form.get(f"group_{st_id}", "").strip()
                sort_val = int(request.form.get(f"order_{st_id}", 99))
                target_val = float(request.form.get(f"target_{st_id}", 3.0))
                is_act = 1 if request.form.get(f"active_{st_id}") else 0
                if st_name:
                    conn.execute("UPDATE stations SET name = ?, group_name = ?, sort_order = ?, target_pi = ?, is_active = ? WHERE id = ?", (st_name, grp_name, sort_val, target_val, is_act, st_id))
            conn.commit()
            flash("Station settings and names updated successfully!", "success")

        elif action == "update_team":
            for eng_id in request.form.getlist("eng_id"):
                formal_name = request.form.get(f"formal_name_{eng_id}", "").strip()
                username = request.form.get(f"username_{eng_id}", "").strip()
                station_id = request.form.get(f"station_{eng_id}")
                station_id = int(station_id) if station_id and station_id.isdigit() else None
                shift = float(request.form.get(f"shift_{eng_id}", 8.0))
                joined_d = request.form.get(f"joined_{eng_id}", "").strip() or None
                act_d = request.form.get(f"activation_{eng_id}", "").strip() or None
                res_d = request.form.get(f"resigned_{eng_id}", "").strip() or None
                transfer_period = request.form.get(f"transfer_{eng_id}", "").strip()
                custom_target_val = request.form.get(f"custom_target_{eng_id}", "").strip()
                custom_target_pi = float(custom_target_val) if custom_target_val else None
                disp_name_val = request.form.get(f"display_name_{eng_id}", "").strip() or None
                track_perf = 1 if request.form.get(f"track_{eng_id}") else 0

                if not username:
                    username = f"admin.{re.sub(r'[^a-z0-9]', '.', formal_name.lower()).strip('.')}"

                cur_eng = conn.execute("SELECT station_id, joined_date, activation_date FROM engineers WHERE id = ?", (eng_id,)).fetchone()
                old_station_id = cur_eng["station_id"] if cur_eng else None
                start_d = parse_iso_date(cur_eng["activation_date"] or cur_eng["joined_date"]) if cur_eng else date(2024, 1, 1)
                init_yr = start_d.year if start_d else 2024
                init_mo = start_d.month if start_d else 1

                conn.execute("""
                    UPDATE engineers 
                    SET formal_name = ?, username = ?, station_id = ?, shift_hours = ?, joined_date = ?, activation_date = ?, resigned_date = ?, custom_target_pi = ?, display_name = ?, track_performance = ?
                    WHERE id = ?
                """, (formal_name, username, station_id, shift, joined_d, act_d, res_d, custom_target_pi, disp_name_val, track_perf, eng_id))

                if transfer_period and station_id and station_id != old_station_id:
                    match = re.match(r'(\d{4})-(\d{2})', transfer_period)
                    if match:
                        t_yr, t_mo = int(match.group(1)), int(match.group(2))
                        if old_station_id:
                            conn.execute("INSERT OR IGNORE INTO engineer_station_history (engineer_id, station_id, effective_year, effective_month) VALUES (?, ?, ?, ?)", (eng_id, old_station_id, init_yr, init_mo))
                        conn.execute("INSERT INTO engineer_station_history (engineer_id, station_id, effective_year, effective_month) VALUES (?, ?, ?, ?) ON CONFLICT(engineer_id, effective_year, effective_month) DO UPDATE SET station_id = excluded.station_id", (eng_id, station_id, t_yr, t_mo))

            conn.commit()
            flash("Active staff details updated successfully!", "success")

        elif action == "update_resigned":
            for eng_id in request.form.getlist("resigned_eng_id"):
                fn = request.form.get(f"res_formal_name_{eng_id}", "").strip()
                u = request.form.get(f"res_username_{eng_id}", "").strip()
                dn = request.form.get(f"res_display_name_{eng_id}", "").strip() or None
                s_id = request.form.get(f"res_station_{eng_id}")
                s_id = int(s_id) if s_id and s_id.isdigit() else None
                jd, ad, rd = request.form.get(f"res_joined_{eng_id}", "").strip() or None, request.form.get(f"res_activation_{eng_id}", "").strip() or None, request.form.get(f"res_resigned_{eng_id}", "").strip() or None
                conn.execute("UPDATE engineers SET formal_name = ?, username = ?, display_name = ?, station_id = ?, joined_date = ?, activation_date = ?, resigned_date = ? WHERE id = ?", (fn, u, dn, s_id, jd, ad, rd, eng_id))

            conn.commit()
            flash("Inactive staff records updated successfully!", "success")

        elif action == "add_station":
            s_name = request.form.get("station_name", "").strip().upper()
            s_order = int(request.form.get("station_order", 99))
            s_target = float(request.form.get("station_target", 3.0))
            if s_name:
                conn.execute("INSERT OR IGNORE INTO stations (name, sort_order, target_pi, is_active) VALUES (?, ?, ?, 1)", (s_name, s_order, s_target))
                conn.commit()
                flash(f"Station '{s_name}' added!", "success")

        elif action == "add_engineer":
            fn = request.form.get("formal_name", "").strip()
            dn = request.form.get("display_name", "").strip() or None
            u = request.form.get("username", "").strip()
            s_id = request.form.get("station_id")
            s_id = int(s_id) if s_id and s_id.isdigit() else None
            shift = float(request.form.get("shift_hours", 8.0))
            jd, ad = request.form.get("joined_date", "").strip() or None, request.form.get("activation_date", "").strip() or None
            custom_target_val = request.form.get("custom_target_pi", "").strip()
            custom_target_pi = float(custom_target_val) if custom_target_val else None
            track_perf = 1 if request.form.get("track_performance") else 0

            if not u: u = f"admin.{re.sub(r'[^a-z0-9]', '.', fn.lower()).strip('.')}"

            if fn:
                try:
                    conn.execute("INSERT INTO engineers (formal_name, display_name, username, station_id, shift_hours, joined_date, activation_date, custom_target_pi, is_active, track_performance) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?)", (fn, dn, u, s_id, shift, jd, ad, custom_target_pi, track_perf))
                    conn.commit()
                    flash(f"Staff member '{fn}' added successfully!", "success")
                except sqlite3.IntegrityError:
                    u_fallback = f"{u}_{int(datetime.now().timestamp())}"
                    conn.execute("INSERT INTO engineers (formal_name, display_name, username, station_id, shift_hours, joined_date, activation_date, custom_target_pi, is_active, track_performance) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?)", (fn, dn, u_fallback, s_id, shift, jd, ad, custom_target_pi, track_perf))
                    conn.commit()
                    flash(f"Staff member '{fn}' added successfully!", "success")

        return redirect(url_for("team"))

    stations = conn.execute("SELECT * FROM stations ORDER BY sort_order ASC, name ASC").fetchall()
    all_engineers = conn.execute("SELECT e.*, s.name as station_name, s.target_pi as default_station_target FROM engineers e LEFT JOIN stations s ON e.station_id = s.id ORDER BY s.sort_order ASC, s.name ASC, e.formal_name ASC").fetchall()
    muted_techs = conn.execute("SELECT * FROM ignored_technicians ORDER BY username").fetchall()
    kpi_configs = conn.execute("SELECT * FROM kpi_configs ORDER BY rowid").fetchall()
    target_hc_row = conn.execute("SELECT setting_val FROM system_settings WHERE setting_key = 'target_headcount'").fetchone()
    target_headcount = int(target_hc_row[0]) if target_hc_row else 23

    today = date.today()
    active_engineers, resigned_engineers = [], []

    for e in all_engineers:
        act_d, res_d = parse_iso_date(e["activation_date"]), parse_iso_date(e["resigned_date"])
        d = dict(e)
        if res_d and res_d <= today:
            d["status"], d["status_class"] = "Resigned", "danger"
            resigned_engineers.append(d)
        elif not e["track_performance"]:
            d["status"], d["status_class"] = "Admin", "info text-dark"
            active_engineers.append(d)
        elif not act_d or act_d > today:
            d["status"], d["status_class"] = "Pending Activation", "warning text-dark"
            active_engineers.append(d)
        else:
            d["status"], d["status_class"] = "Active", "success"
            active_engineers.append(d)

    conn.close()
    return render_template("team.html", stations=stations, active_engineers=active_engineers, resigned_engineers=resigned_engineers, muted_techs=muted_techs, kpi_configs=kpi_configs, target_headcount=target_headcount)

@app.route("/team/reactivate/<int:eng_id>", methods=["POST"])
def reactivate_engineer(eng_id):
    conn = get_db()
    conn.execute("UPDATE engineers SET resigned_date = NULL, is_active = 1 WHERE id = ?", (eng_id,))
    conn.commit()
    conn.close()
    flash("Engineer restored to active roster!", "success")
    return redirect(url_for("team"))


# =============================================================================
# GITHUB AUTO-DEPLOY WEBHOOK
# =============================================================================
@app.route("/api/github-deploy", methods=["POST"], strict_slashes=False)
def github_deploy():
    import subprocess

    # 1. Pull latest code from GitHub
    repo_dir = os.path.dirname(os.path.abspath(__file__))
    try:
        subprocess.run(["git", "-C", repo_dir, "pull"], check=True)
    except Exception as e:
        return f"Git pull failed: {e}", 500

    # 2. Touch WSGI file to trigger PythonAnywhere instant reload
    wsgi_file = os.path.expanduser("~/var/www/_wsgi_reload.txt")
    for f in os.listdir("/var/www/"):
        if f.endswith("_wsgi.py"):
            os.utime(os.path.join("/var/www/", f), None)
            break

    return "Code pulled and web app reloaded successfully!", 200

if __name__ == "__main__":
    app.run(debug=True, port=5000)